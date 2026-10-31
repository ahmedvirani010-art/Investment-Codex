"""Remaining PRD: evaluation, lineage, migration, cache and wealth failure gates."""

import json
from copy import deepcopy
from datetime import date
from decimal import Decimal
from importlib.resources import files

import pytest
from openpyxl import Workbook, load_workbook

from investment_workspace import pipeline, storage
from investment_workspace.cli import main
from investment_workspace.intake import import_package, migrate_workbook
from investment_workspace.models import REGISTRY, reference_value, bridge
from investment_workspace.readiness import committed_report, workflow_status
from investment_workspace.schemas import Bundle, WorkbookMigration, ValuationModel, Scenario
from investment_workspace.storage import dumps, digest, read_json


@pytest.fixture
def full():
    return json.loads(files('investment_workspace').joinpath('fixtures/phase4-demo.json').read_bytes())


def run(repo, data, **options):
    result=pipeline.run(repo,dumps(data),**options)
    return result,read_json(repo,pipeline.STATE)


def asset_plan(full, balance='1200', months=12, spending='0', retirement='2028-01-01'):
    w=full['wealth']
    w['assets']=[w['assets'][0]]
    w['assets'][0]['balance']=balance
    w['flows']=[];w['histories']=[]
    p=w['plans'][0];p.update(months=months,monthly_spending_real=spending,retirement_date=retirement,withdrawal_order=['EXTERNAL_CASH'])
    for s in w['scenarios']:
        s['annual_inflation']='0';s['returns']=[s['returns'][0]]
        s['returns'][0].update(annual_nominal_return='0',annual_fee_rate='0',return_tax_rate='0',withdrawal_tax_rate='0')
    return w


def test_evaluation_date_expires_old_inputs_without_relabelling(repo,full):
    _,before=run(repo,full)
    result,after=run(repo,full,as_of='2026-11-10',workflow='daily')
    assert result['status']=='PARTIAL'
    assert after['summary']['market_cutoff']=='2026-10-09'
    assert after['evaluation_as_of']=='2026-11-10'
    assert all(r['freshness']=='STALE' for r in after['holdings'])
    assert after['holdings'][0]['price_date']==before['holdings'][0]['price_date']
    assert all(r['rs_percentile'] is None for r in after['integration']['timing'])
    assert not after['research']['ranking']
    assert after['wealth']['summary']['retirement_status']=='PARTIAL'
    assert pipeline.validate(repo)['workbook_parity']=='PASS'


def test_evaluation_cannot_backdate_observed_inputs(repo,full):
    result,state=run(repo,full,as_of='2026-10-08')
    assert result['status']=='FAILED'
    assert state['holdings']==[]


def test_owner_specific_completeness_is_not_legacy_migration_status(repo,full):
    _,state=run(repo,full)
    assert workflow_status(state,'timing')=='COMPLETE'
    assert workflow_status(state,'retirement')=='COMPLETE'
    assert workflow_status(state,'daily')=='PARTIAL'
    assert committed_report(repo)['compatibility']=='NOT_MIGRATED'
    assert main(['--root',str(repo),'run','timing','--require-full-coverage'])==0
    assert main(['--root',str(repo),'run','daily','--require-full-coverage'])==1


def test_cache_reuse_and_bounded_resume(repo,full):
    first,partial=run(repo,full,max_units=2)
    assert first['status']=='PARTIAL'
    assert partial['deferred_work']
    assert len(partial['calculations'])==2
    _,complete=run(repo,full)
    assert len(complete['calculations'])>2
    assert complete['usage'][-1]['cache_hits']>=2
    again,unchanged=run(repo,full)
    assert again['status']=='NO_MATERIAL_CHANGE'
    assert unchanged['usage'][-1]['calculation_units']==0
    assert unchanged['usage'][-1]['cache_hits']>0
    assert unchanged['wealth']==complete['wealth']
    assert unchanged['usage'][-1]['input_tokens'] is None


def test_changed_return_invalidates_only_related_projection_cache(repo,full):
    _,before=run(repo,full)
    scenario=full['wealth']['scenarios'][1]
    old_id=scenario['record_id']
    new=deepcopy(full['sources'][-1]);new['source_id']='synthetic.wealth.newpolicy';full['sources'].append(new)
    scenario.update(record_id='new.base.2',supersedes=old_id,source_ids=[new['source_id']])
    scenario['returns'][0]['annual_nominal_return']='0.15'
    _,after=run(repo,full)
    assert after['usage'][-1]['calculation_units']==1
    assert after['research']['valuations']==before['research']['valuations']
    assert after['wealth']['projections'][1]['ending_nominal_balance']!=before['wealth']['projections'][1]['ending_nominal_balance']


def test_wealth_hand_calculated_contributions_month_end(repo,full):
    w=asset_plan(full)
    flow=deepcopy(full['wealth']['plans'][0])
    # Reuse only the shared record fields for a strictly typed net contribution.
    flow={k:flow[k] for k in ('ticker','effective_date','valid_until','source_ids')}
    w['flows']=[{**flow,'record_id':'contribution.1','asset_key':'EXTERNAL_CASH','kind':'CONTRIBUTION','amount':'100',
                 'amount_basis':'NET','price_basis':'NOMINAL','frequency':'MONTHLY','starts':'2026-11-01','ends':'2027-10-31'}]
    result,state=run(repo,full,workflow='retirement')
    assert result['status']=='COMPLETE'
    assert all(Decimal(r['ending_nominal_balance'])==2400 for r in state['wealth']['projections'])
    assert all(Decimal(r['ending_real_balance'])==2400 for r in state['wealth']['projections'])
    assert len(state['wealth']['retirement'])==36
    assert state['summary']['cash']=='15220.00'
    assert pipeline.validate(repo)['workbook_parity']=='PASS'


def test_inflation_and_geometric_monthly_return(repo,full):
    w=asset_plan(full,balance='1000')
    for scenario in w['scenarios']:
        scenario['annual_inflation']='0.12'
        scenario['returns'][0]['annual_nominal_return']='0.12'
    _,state=run(repo,full)
    for summary in state['wealth']['projections']:
        assert abs(Decimal(summary['ending_nominal_balance'])-1120)<Decimal('1e-20')
        assert abs(Decimal(summary['ending_real_balance'])-1000)<Decimal('1e-20')


def test_locked_assets_and_shortfalls_are_not_borrowing(repo,full):
    w=full['wealth'];w['assets']=w['assets'][:2];w['assets'][0]['balance']='100';w['flows']=[];w['histories']=[]
    p=w['plans'][0];p.update(months=3,monthly_spending_real='100',retirement_date='2026-11-01',withdrawal_order=['EXTERNAL_CASH','PF'])
    for scenario in w['scenarios']:
        scenario['annual_inflation']='0';scenario['returns']=scenario['returns'][:2]
        for policy in scenario['returns']:
            policy['annual_nominal_return']='0'
    _,state=run(repo,full)
    base=[r for r in state['wealth']['retirement'] if r['scenario']=='base']
    assert [Decimal(r['funding_gap']) for r in base]==[0,100,0]
    assert Decimal(base[1]['locked_balance'])==600
    assert Decimal(base[2]['nominal_balance'])==500
    assert all(Decimal(v)>=0 for r in base for v in r['asset_balances'].values())
    summary=next(r for r in state['wealth']['projections'] if r['scenario']=='base')
    assert summary['first_depletion_date']=='2026-12-31'
    assert summary['borrowing_assumed'] is False


def test_explicit_tax_net_withdrawal_and_unknown_tax(repo,full):
    w=asset_plan(full,balance='1000',months=1,spending='90',retirement='2026-11-01')
    for s in w['scenarios']:
        s['returns'][0]['withdrawal_tax_rate']='0.1'
    _,state=run(repo,full)
    row=next(r for r in state['wealth']['retirement'] if r['scenario']=='base')
    assert Decimal(row['nominal_balance'])==900
    assert Decimal(row['net_withdrawn'])==90
    assert Decimal(row['taxes'])==10
    # A fresh dataset with unknown tax must not manufacture a zero-tax projection.


def test_unknown_returns_taxes_and_balances_block_projection(repo,full):
    w=asset_plan(full)
    w['scenarios'][0]['returns'][0]['return_tax_rate']=None
    result,state=run(repo,full,workflow='retirement')
    assert result['status']=='PARTIAL'
    bear=next(r for r in state['wealth']['projections'] if r['scenario']=='bear')
    assert bear['ending_nominal_balance'] is None
    assert not any(r['scenario']=='bear' for r in state['wealth']['retirement'])


def test_portfolio_links_cannot_double_count_cash_or_benefits(full):
    a=full['wealth']['assets'][0]
    a.update(portfolio_link='TOTAL',balance=None)
    b=full['wealth']['assets'][1];b.update(portfolio_link='CASH',balance=None)
    with pytest.raises(ValueError,match='counted twice'):
        Bundle.model_validate(full)
    b.pop('portfolio_link');b['balance']='600'
    a.pop('portfolio_link');a['balance']='100'
    flow=full['wealth']['flows'][0]
    flow.update(kind='BENEFIT',economic_claim_id=b['economic_id'])
    with pytest.raises(ValueError,match='claim already'):
        Bundle.model_validate(full)


def test_crypto_opt_in_and_missing_fx_stay_explicit(repo,full):
    full['wealth']['assets'][2]['kind']='CRYPTO'
    with pytest.raises(ValueError,match='opt-in'):
        Bundle.model_validate(full)
    full['wealth']['include_crypto']=True
    full['wealth']['assets'][2]['currency']='USD'
    _,state=run(repo,full)
    assert state['wealth']['summary']['total_wealth_pkr'] is None
    assert state['wealth']['summary']['retirement_status']=='PARTIAL'


def test_cross_asset_fx_and_zero_variance(repo,full):
    h=full['wealth']['histories'][1]
    h['currency']='USD';h['fx_source_id']='synthetic.wealth.history'
    for i,q in enumerate(h['quotes']):
        q['fx_to_pkr']=str(100+i)
    _,state=run(repo,full)
    row=next(r for r in state['wealth']['cross_asset'] if r['asset_key']=='GOLD')
    expected=Decimal(105)*103/(Decimal(100)*100)-1
    assert Decimal(row['holding_period_return'])==expected
    assert state['wealth']['correlations'][0]['correlation'] is None
    assert state['wealth']['correlations'][0]['status']=='UNDEFINED_ZERO_VARIANCE'
    assert state['wealth']['summary']['cross_asset_status']=='COMPLETE'


def test_partial_history_is_not_zero_return(repo,full):
    full['wealth']['histories'][1]['quotes'][1]['close']=None
    _,state=run(repo,full)
    row=next(r for r in state['wealth']['cross_asset'] if r['asset_key']=='GOLD')
    assert row['holding_period_return'] is None
    assert state['wealth']['summary']['cross_asset_status']=='PARTIAL'


def legacy_files(tmp_path,state, *, formula=False):
    wb=Workbook();ws=wb.active;ws.title='Legacy RS';ws.append(['Symbol','RS percentile','Long SMA'])
    for row in state['integration']['timing']:
        ws.append([row['ticker'],float(row['rs_percentile']),float(row['long_sma'])])
    history=wb.create_sheet('Legacy Prices');history.append(['Symbol','Session','Close'])
    for record in state['integration']['input']['histories']:
        if record.get('role','SECURITY')!='SECURITY': continue
        for bar in record['bars']: history.append([record['ticker'],bar['session_date'],float(bar['close'])])
    if formula: ws['B2']='=50+50' 
    workbook=tmp_path/'legacy.xlsx';wb.save(workbook);wb.close()
    sha=digest(workbook.read_bytes())
    migration={'record_id':'legacy.migration.1','ticker':'DATASTORE','effective_date':'2026-10-09','valid_until':'2026-10-31',
       'source_ids':['legacy.workbook.1'],'kind':'LEGACY_DATASTORE','workbook_sha256':sha,'approved_by':'synthetic tester',
       'reason':'SYNTHETIC mapping parity fixture','store_permitted':True,'methodology_reviewed':True,
       'mappings':[{'sheet':'Legacy RS','collection':'timing','columns':{'ticker':'Symbol','rs_percentile':'RS percentile','long_sma':'Long SMA'}},
                   {'sheet':'Legacy Prices','collection':'history','key_fields':['ticker','session_date'],'columns':{'ticker':'Symbol','session_date':'Session','close':'Close'}}]}
    source={'source_id':'legacy.workbook.1','title':'SYNTHETIC original datastore','publisher':'test fixture','effective_date':'2026-10-09',
            'locator':'Legacy RS','evidence_label':'CONFIRMED','document_sha256':sha}
    mapping=tmp_path/'mapping.json';mapping.write_bytes(dumps({'migration':migration,'sources':[source]}))
    return workbook,mapping


def test_migration_dry_run_accepts_only_actual_numerical_parity(repo,full,tmp_path):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state)
    old_head=storage.head(repo)
    preview=migrate_workbook(repo,workbook,mapping)
    assert preview['status']=='VERIFIED' and preview['accepted'] is False
    assert storage.head(repo)==old_head
    result=migrate_workbook(repo,workbook,mapping,accept=True)
    assert result['status']!='FAILED', result.get('error')
    accepted=read_json(repo,pipeline.STATE)
    assert accepted['integration']['summary']['legacy_datastore_status']=='VERIFIED'
    assert not any(g['reason']=='LEGACY_DATASTORE_NOT_MIGRATED' for g in accepted['data_gaps'])
    assert storage.committed(repo,'workbooks/PSX_Mean_Reversion_DataStore.xlsx')
    assert storage.committed(repo,'data/evidence/blobs/'+digest(workbook.read_bytes())+'.xlsx')==workbook.read_bytes()
    assert pipeline.validate(repo)['integrity']=='PASS'


def test_missing_formula_cache_and_numeric_mismatch_block_migration(repo,full,tmp_path):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state,formula=True)
    with pytest.raises(ValueError,match='cached result'):
        migrate_workbook(repo,workbook,mapping)
    workbook,mapping=legacy_files(tmp_path,state)
    wb=load_workbook(workbook);wb.active['B2']=1;wb.save(workbook);wb.close()
    data=json.loads(mapping.read_bytes());data['migration']['workbook_sha256']=digest(workbook.read_bytes())
    mapping.write_bytes(dumps(data))
    with pytest.raises(ValueError,match='parity mismatch'):
        migrate_workbook(repo,workbook,mapping)


def test_package_archives_licensed_document_and_rejects_changed_hash(repo,full,tmp_path):
    run(repo,full)
    doc=tmp_path/'contract.md';doc.write_text('SYNTHETIC evidence contract fixture')
    source={'source_id':'contract.test.1','title':'SYNTHETIC contract','publisher':'test','effective_date':'2026-10-09',
            'locator':'contract.md','evidence_label':'CONFIRMED','document_sha256':digest(doc.read_bytes()),
            'accessed_at':'2026-10-26T00:00:00Z','coverage_through':'2026-10-09','owner':'evidence'}
    package={'sources':[source],'files':[{'path':'contract.md','document':{'source_id':source['source_id'],'sha256':digest(doc.read_bytes()),
              'suffix':'.md','role':'EVIDENCE_CONTRACT','store_permitted':True}}]}
    path=tmp_path/'packet.json';path.write_bytes(dumps(package))
    assert import_package(repo,path)['status']!='FAILED'
    state=read_json(repo,pipeline.STATE)
    assert state['documents'][0]['role']=='EVIDENCE_CONTRACT'
    assert storage.committed(repo,'data/evidence/blobs/'+digest(doc.read_bytes())+'.md')==doc.read_bytes()
    previous=deepcopy(state['wealth'])
    doc.write_text('changed after declared hash')
    assert import_package(repo,path)['status']=='FAILED'
    assert read_json(repo,pipeline.STATE)['wealth']==previous


def real_bundle():
    return {'dataset_id':'production','mode':'USER_IMPORT','market_cutoff':'2026-10-09',
       'sources':[{'source_id':'user.balances.1','title':'User-confirmed opening and closing','publisher':'user','effective_date':'2026-10-09','locator':'statement','evidence_label':'CONFIRMED'}],
       'securities':[], 'prices':[], 'transactions':[],
       'opening':{'timestamp':'2026-10-08T18:00:00+05:00','source_id':'user.balances.1','cash':'1000','positions':[]},
       'closing':{'timestamp':'2026-10-09T18:00:00+05:00','source_id':'user.balances.1','cash':'1000','positions':[]}}


def test_dataset_boundary_archives_demo_and_never_mixes_ledger(repo,full):
    _,old=run(repo,full)
    real=real_bundle()
    bad,state=run(repo,real)
    assert bad['status']=='FAILED'
    assert state['transactions']==old['transactions']
    result,active=run(repo,real,activate_dataset={'author':'test user','reason':'controlled synthetic boundary test','confirmed_real_data':True})
    assert result['status']=='SUCCESS'
    assert active['transactions']==[]
    assert active['dataset']['dataset_id']=='production'
    path=f'data/history/datasets/demo.full/{state["published_run_id"]}.json'
    assert storage.committed(repo,path)
    assert len(active['run_log'])==3


def test_a_mode_flip_cannot_promote_fixture_balances(repo,full):
    run(repo,full)
    full['mode']='USER_IMPORT';full['dataset_id']='production'
    result,_=run(repo,full,activate_dataset={'author':'test','reason':'attempted fixture promotion','confirmed_real_data':True})
    assert result['status']=='FAILED'
    assert 'synthetic fixtures' in result['error']


def test_wealth_locked_publication_recovery_preserves_context(repo,full,monkeypatch):
    actual=storage.atomic_write
    def locked(path,data):
        if path==repo/pipeline.WORKBOOK: raise PermissionError('locked workbook')
        return actual(path,data)
    monkeypatch.setattr(storage,'atomic_write',locked)
    with pytest.raises(storage.PublicationError):
        run(repo,full,as_of='2026-10-10',workflow='retirement')
    monkeypatch.setattr(storage,'atomic_write',actual)
    result=pipeline.recover(repo)
    assert result['status']!='FAILED'
    state=read_json(repo,pipeline.STATE)
    assert state['evaluation_as_of']=='2026-10-10'
    assert state['wealth']['summary']['retirement_status']=='COMPLETE'
    assert pipeline.validate(repo)['workbook_parity']=='PASS'


def test_automation_defaults_cannot_start_or_publish(repo,full):
    from investment_workspace.automation import execute
    run(repo,full)
    old=storage.head(repo)
    result=execute(repo,'daily','2026-10-26',dry_run=False,scheduled=True)
    assert result['status']=='BLOCKED'
    assert result['mutated'] is False
    assert storage.head(repo)==old


def test_current_cells_alone_cannot_certify_history(repo,full,tmp_path):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state)
    wb=load_workbook(workbook);del wb['Legacy Prices'];wb.save(workbook);wb.close()
    data=json.loads(mapping.read_bytes())
    data['migration']['mappings']=data['migration']['mappings'][:1]
    sha=digest(workbook.read_bytes());data['migration']['workbook_sha256']=sha;data['sources'][0]['document_sha256']=sha
    mapping.write_bytes(dumps(data))
    assert migrate_workbook(repo,workbook,mapping)['status']=='PARTIAL_MAPPING'
    result=migrate_workbook(repo,workbook,mapping,accept=True)
    assert result['status']=='PARTIAL'
    after=read_json(repo,pipeline.STATE)
    assert after['integration']['summary']['legacy_datastore_status']=='NOT_MIGRATED'
    assert storage.committed(repo,'workbooks/PSX_Mean_Reversion_DataStore.xlsx') is None


def test_unmapped_sheet_needs_explicit_review_reason(repo,full,tmp_path):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state)
    wb=load_workbook(workbook);wb.create_sheet('Additional evidence').append(['notes']);wb.save(workbook);wb.close()
    data=json.loads(mapping.read_bytes());sha=digest(workbook.read_bytes());data['migration']['workbook_sha256']=sha
    mapping.write_bytes(dumps(data))
    with pytest.raises(ValueError,match='unmapped legacy sheets'):
        migrate_workbook(repo,workbook,mapping)


def test_verified_compatibility_stale_refresh_exports_unknown_not_old_numbers(repo,full,tmp_path):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state)
    accepted=migrate_workbook(repo,workbook,mapping,accept=True)
    assert accepted['status']!='FAILED'
    result=pipeline.run(repo,pipeline.current_input(repo),as_of='2026-11-10',workflow='daily')
    assert result['status']=='PARTIAL'
    wb=load_workbook(repo/'workbooks/PSX_Mean_Reversion_DataStore.xlsx',data_only=True)
    assert wb['Legacy RS']['B2'].value=='UNKNOWN'
    metadata=dict(list(wb['_Workspace_Provenance'].values)[1:])
    assert metadata['evaluation_as_of']=='2026-11-10'
    wb.close()
    assert pipeline.validate(repo)['integrity']=='PASS'


def test_recovery_retains_original_artifact_bytes(repo,full,tmp_path,monkeypatch):
    _,state=run(repo,full)
    workbook,mapping=legacy_files(tmp_path,state)
    original=storage.atomic_write
    def locked(path,data):
        if path==repo/pipeline.WORKBOOK: raise PermissionError('locked')
        return original(path,data)
    monkeypatch.setattr(storage,'atomic_write',locked)
    with pytest.raises(storage.PublicationError):
        migrate_workbook(repo,workbook,mapping,accept=True)
    monkeypatch.setattr(storage,'atomic_write',original)
    pipeline.recover(repo)
    sha=digest(workbook.read_bytes())
    assert storage.committed(repo,'data/evidence/blobs/'+sha+'.xlsx')==workbook.read_bytes()
    assert read_json(repo,pipeline.STATE)['integration']['summary']['legacy_datastore_status']=='VERIFIED'
    assert pipeline.validate(repo)['integrity']=='PASS'


def test_wealth_cell_tampering_is_detected(repo,full,tmp_path):
    from investment_workspace.excel import validate_workbook
    _,state=run(repo,full)
    wb=load_workbook(repo/pipeline.WORKBOOK)
    ws=wb['Retirement'];headers=[c.value for c in ws[1]]
    ws.cell(2,headers.index('nominal_balance')+1,0)
    target=tmp_path/'tampered.xlsx';wb.save(target);wb.close()
    with pytest.raises(ValueError,match='parity'):
        validate_workbook(target,state)


def test_evidence_change_is_visible_without_fabricated_trade(repo,full,tmp_path):
    _,before=run(repo,full)
    source=deepcopy(full['sources'][0]);source['source_id']='new.corrob.source';full['sources'].append(source)
    _,after=run(repo,full)
    assert after['transactions']==before['transactions']
    assert any(c.get('change_type')=='EVIDENCE' for c in after['changes'] if c['run_id']==after['published_run_id'])
    report=pipeline.changes_since(repo,before['published_run_id'])
    assert any(c['field']=='evidence_registry' for c in report['changes'])


def test_automation_checks_no_policy_even_for_manual_dry_run(repo,full):
    run(repo,full)
    old=storage.head(repo)
    assert main(['--root',str(repo),'automation','--workflow','daily','--as-of','2026-10-26','--dry-run'])==1
    assert storage.head(repo)==old


def test_wealth_cli_import_and_replay(repo,full,tmp_path):
    wealth=full.pop('wealth')
    used={s for group in ('assets','flows','plans','scenarios','histories') for r in wealth[group] for s in r['source_ids']}
    sources=[s for s in full['sources'] if s['source_id'] in used]
    full['sources']=[s for s in full['sources'] if s['source_id'] not in used]
    run(repo,full)
    packet=tmp_path/'wealth.json';packet.write_bytes(dumps({'sources':sources,'wealth':wealth}))
    assert main(['--root',str(repo),'import-wealth','--file',str(packet)])==0
    state=read_json(repo,pipeline.STATE)
    assert main(['--root',str(repo),'run','retirement','--require-full-coverage'])==0
    target=tmp_path/'replayed.xlsx'
    assert pipeline.replay(repo,state['published_run_id'],target)['workbook_parity']=='PASS'


def test_linked_portfolio_is_referenced_once_and_keeps_psx_totals_separate(repo,full):
    _,state=run(repo,full)
    row=next(a for a in state['wealth']['assets'] if a['ticker']=='PSX_TOTAL')
    assert Decimal(row['balance_pkr'])==Decimal(state['summary']['total_assets'])
    assert Decimal(state['wealth']['summary']['total_wealth_pkr'])==Decimal(state['summary']['total_assets'])+2100
    assert state['summary']['cash']=='15220.00'
    assert state['wealth']['summary']['portfolio_included'] is True


def test_long_excel_fields_have_deterministic_canonical_references():
    from investment_workspace.excel import cell_value
    large=[{'item':i,'value':'a'*100} for i in range(1000)]
    value=cell_value(large)
    assert len(value)<1000
    assert json.loads(value)['status']=='FULL_VALUE_IN_CANONICAL_SNAPSHOT'
    assert cell_value('1e-1000','price')=='1e-1000'
    assert cell_value('1e1000','price')=='1e1000'


def test_current_valuation_requires_current_event_review(repo, full):
    _, state = run(repo, full, as_of='2026-10-10')
    assert not state['research']['ranking']
    assert any(row['outcome'] == 'BLOCKED_EVIDENCE' for row in state['coverage'])


def test_event_review_cannot_claim_future_coverage(full):
    full['research']['models'][0]['event_checked_through'] = '2026-10-10'
    with pytest.raises(ValueError, match='future event coverage'):
        Bundle.model_validate(full)
