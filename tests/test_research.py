"""Hand-calculated synthetic reference models and Phase 2 failure gates."""

from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from investment_workspace import pipeline, storage
from investment_workspace.cli import main
from investment_workspace.schemas import Bundle
from investment_workspace.storage import dumps, read_json


@pytest.fixture
def research_bundle():
    import json
    from importlib.resources import files
    return json.loads(files("investment_workspace").joinpath("fixtures/phase2-demo.json").read_bytes())


def run(repo, bundle):
    result = pipeline.run(repo, dumps(bundle), workflow="valuation")
    return result, read_json(repo, pipeline.STATE)


def test_reference_models_bridge_returns_and_coverage(repo, research_bundle):
    result, state = run(repo, research_bundle)
    assert result["status"] == "PARTIAL"  # Unsupported loss sector cannot become a generic P/E model.
    research = state["research"]
    bank = next(r for r in research["valuations"] if r["ticker"] == "DEMO-BANK" and r["scenario"] == "base")
    assert Decimal(bank["reported_eps"]) == 10
    assert Decimal(bank["normalized_eps"]) == Decimal("9.3")  # (1000 - 100 + 30) / 100
    assert Decimal(bank["forward_eps"]) == 12
    assert Decimal(bank["present_fair_value"]) == 225  # 100*(.20-.02)/(.10-.02)
    assert Decimal(bank["gross_return"]) == Decimal("0.04")  # (125 + 5 - 125)/125, not (FV-price)/price
    assert bank["net_return"] is None
    assert bank["assumption_ids"][0] in {a["assumption_id"] for a in state["assumptions"]}
    industrial = next(r for r in research["valuations"] if r["ticker"] == "DEMO-INDUSTRIAL")
    assert abs(Decimal(industrial["present_fair_value"]) - Decimal("118.1818181818181818")) < Decimal("1e-14")
    assert research["summary"]["eligible_count"] == 3
    assert len(research["ranking"]) == 3
    assert all(r["ticker"] != "DEMO-LOSS" for r in research["ranking"])
    entry = next(r for r in research["entries"] if r["ticker"] == "DEMO-BANK")
    assert Decimal(entry["mos_entry_ceiling"]) == 180
    assert abs(Decimal(entry["entry_ceiling"]) - Decimal(130)/Decimal("1.15")) < Decimal("1e-25")
    assert entry["entry_status"] == "BLOCKED"
    assert state["summary"]["sizing_status"] == "BLOCKED"
    pipeline.validate(repo)


@pytest.mark.parametrize("change,expected", [
    ("missing_earnings", "BLOCKED_EVIDENCE"), ("stale", "BLOCKED_EVIDENCE"),
    ("provisional", "BLOCKED_MODEL"), ("legacy", "BLOCKED_MODEL"),
    ("basis", "BLOCKED_EVIDENCE"), ("event", "BLOCKED_EVIDENCE"),
])
def test_model_gates(repo, research_bundle, change, expected):
    r = research_bundle["research"]
    m = r["models"][0]
    if change == "missing_earnings":
        r["earnings"].pop(0)
    elif change == "stale":
        m["valid_until"] = "2026-10-09"
        research_bundle["market_cutoff"] = "2026-10-10"
    elif change == "provisional":
        m["certification"] = "PROVISIONAL"
    elif change == "legacy":
        m["origin"] = "LEGACY_RECOVERED"
    elif change == "basis":
        m["accounting_basis"] = "CONSOLIDATED"
    else:
        m["event_checked_through"] = None
    _, state = run(repo, research_bundle)
    assert next(c for c in state["coverage"] if c["ticker"] == "DEMO-BANK")["outcome"] == expected
    assert all(r["ticker"] != "DEMO-BANK" for r in state["research"]["ranking"])


def test_no_change_preserves_lineage_and_replay(repo, research_bundle, tmp_path):
    first, state = run(repo, research_bundle)
    second, after = run(repo, research_bundle)
    assert second["status"] == "NO_MATERIAL_CHANGE"
    assert state["research"] == after["research"]
    assert after["analysis_run_id"] == first["run_id"]
    pipeline.replay(repo, first["run_id"], tmp_path / "research.xlsx")
    wb = load_workbook(tmp_path / "research.xlsx")
    assert len(list(wb["Valuation"].values)) == 10
    wb.close()


def test_immutable_revision_and_omission_fail_keep_last_good(repo, research_bundle):
    _, initial = run(repo, research_bundle)
    research_bundle["research"]["models"][0]["scenarios"][0]["forward_eps"] = "99"
    failed, after = run(repo, research_bundle)
    assert failed["status"] == "FAILED"
    assert initial["research"] == after["research"]
    research_bundle["research"]["models"][0]["scenarios"][0]["forward_eps"] = "12"
    research_bundle["research"]["models"].pop(0)
    assert run(repo, research_bundle)[0]["status"] == "FAILED"


def dividend(ticker, record_id):
    return {"record_id": record_id, "ticker": ticker, "effective_date": "2026-10-09",
        "valid_until": "2026-10-12", "source_ids": ["synthetic.research.v1"],
        "announcement_date": "2026-10-08", "dps": "2.5", "ex_date": "2026-10-09",
        "record_date": "2026-10-09", "payment_date": "2026-10-09", "dates_verified": True,
        "settlement_rule": "SETTLED_BEFORE_EX", "tax_rate": "0.2"}


def test_entitlement_before_ex_and_cash_not_double_counted(repo, research_bundle):
    bank = dividend("DEMO-BANK", "div.bank")
    bank["receipt_transaction_ids"] = ["demo.receipt.1"]
    energy = dividend("DEMO-ENERGY", "div.energy")
    research_bundle["research"]["dividends"] = [bank, energy]
    _, state = run(repo, research_bundle)
    bank, energy = state["research"]["dividends"]
    assert bank["eligible_quantity"] == "100"  # Sale on ex-date retains earlier entitlement.
    assert bank["gross_entitlement"] == "250.0"
    assert bank["status"] == "RECEIVED"
    assert Decimal(bank["net_received"]) == 200
    assert bank["gross_pending"] == "0"
    assert energy["eligible_quantity"] == "0"  # Buy on ex-date gets no entitlement.
    assert energy["status"] == "NOT_ENTITLED"
    assert state["summary"]["cash"] == "15220.00"


def test_unverified_dates_and_old_window_remain_unknown(repo, research_bundle):
    d = dividend("DEMO-BANK", "div.bank")
    d["dates_verified"] = False
    research_bundle["research"]["dividends"] = [d]
    _, state = run(repo, research_bundle)
    assert state["research"]["dividends"][0]["eligible_quantity"] is None
    assert state["research"]["dividends"][0]["status"] == "PENDING_CONFIRMATION"


def test_duplicate_receipt_link_fails(repo, research_bundle):
    d = dividend("DEMO-BANK", "div.bank")
    d["receipt_transaction_ids"] = ["demo.receipt.1", "demo.receipt.1"]
    research_bundle["research"]["dividends"] = [d]
    assert run(repo, research_bundle)[0]["status"] == "FAILED"


def test_complete_requires_full_universe_and_all_eligible(repo, research_bundle, tmp_path):
    # Add a deliberately explicit adapter-compatible synthetic sector for this fixture.
    research_bundle["securities"][-1]["sector"] = "Industrial"
    result, state = run(repo, research_bundle)
    assert result["status"] == "COMPLETE"
    assert state["research"]["summary"]["valuation_status"] == "COMPLETE"
    assert main(["--root", str(repo), "run", "valuation", "--offline", "--scope", "required", "--require-full-coverage"]) == 0
    packet = {"sources": [], "research": deepcopy(research_bundle["research"])}
    packet["research"]["full_universe_verified"] = False
    path = tmp_path / "packet.json"
    path.write_bytes(dumps(packet))
    # Same universe source cannot be silently changed.
    assert main(["--root", str(repo), "import-research", "--file", str(path)]) == 1


def test_negative_eps_does_not_prevent_cashflow_model(repo, research_bundle):
    research_bundle["securities"][-1]["sector"] = "Industrial"
    _, state = run(repo, research_bundle)
    assert any(r["ticker"] == "DEMO-LOSS" for r in state["research"]["valuations"])


def test_invalid_growth_and_recurring_adjustments_rejected(research_bundle):
    research_bundle["research"]["models"][0]["scenarios"][0]["terminal_growth"] = "0.1"
    with pytest.raises(ValueError, match="terminal growth"):
        Bundle.model_validate(research_bundle)
    research_bundle["research"]["models"][0]["scenarios"][0]["terminal_growth"] = "0.02"
    research_bundle["research"]["earnings"][0]["adjustments"][0]["recurring_test"] = "RECURRING"
    with pytest.raises(ValueError):
        Bundle.model_validate(research_bundle)


def test_phase2_locked_workbook_recovery(repo, research_bundle, monkeypatch):
    original = storage.atomic_write
    def locked(path, data):
        if path == repo / pipeline.WORKBOOK:
            raise PermissionError("Excel locked")
        return original(path, data)
    monkeypatch.setattr(storage, "atomic_write", locked)
    with pytest.raises(storage.PublicationError):
        run(repo, research_bundle)
    monkeypatch.setattr(storage, "atomic_write", original)
    pipeline.recover(repo)
    assert read_json(repo, pipeline.STATE)["research"]["summary"]["eligible_count"] == 3
    pipeline.validate(repo)


def test_explicit_supersession_retains_history(repo, research_bundle):
    _, before = run(repo, research_bundle)
    old = research_bundle["research"]["models"][0]
    new = deepcopy(old)
    new.update(record_id="bank.model.2", supersedes=old["record_id"], source_ids=["synthetic.revision.2"])
    new["scenarios"][1]["forward_eps"] = "13"
    research_bundle["research"]["models"][0] = new
    source = deepcopy(research_bundle["sources"][-1])
    source["source_id"] = "synthetic.revision.2"
    research_bundle["sources"].append(source)
    result, after = run(repo, research_bundle)
    assert result["status"] == "PARTIAL"
    assert len(after["research"]["history"]) == len(before["research"]["history"]) + 1
    assert after["research"]["history"][:-1] == before["research"]["history"]
    assert any(r["record_id"] == "bank.model.2" and r["forward_eps"] == "13" for r in after["research"]["valuations"])


def test_legacy_conflict_requires_evidenced_resolution(repo, research_bundle):
    # Rename a synthetic case to test the hard migration gate, not a real POL valuation.
    import json
    fixture = json.loads(json.dumps(research_bundle).replace("DEMO-BANK", "POL"))
    _, blocked = run(repo, fixture)
    assert next(c for c in blocked["coverage"] if c["ticker"] == "POL")["reason"] == "Unresolved legacy conflict: POL_DPS"
    source = deepcopy(fixture["sources"][-1])
    source["source_id"] = "synthetic.resolution.1"
    fixture["sources"].append(source)
    fixture["research"]["resolutions"] = [{"record_id": "POL.resolution", "ticker": "POL",
        "effective_date": "2026-10-09", "valid_until": "2026-10-12", "source_ids": [source["source_id"]],
        "issue": "POL_DPS", "rationale": "Synthetic conflict resolution test only"}]
    result, resolved = run(repo, fixture)
    assert result["status"] == "PARTIAL"
    assert any(r["ticker"] == "POL" for r in resolved["research"]["valuations"])


def test_split_before_ex_date_and_receipt_mismatch(repo, research_bundle):
    research_bundle["market_cutoff"] = "2026-10-10"
    research_bundle["closing"]["timestamp"] = "2026-10-10T18:00:00+05:00"
    d = dividend("DEMO-INDUSTRIAL", "div.industrial")
    d.update(ex_date="2026-10-10", record_date="2026-10-10", payment_date="2026-10-10")
    research_bundle["research"]["dividends"] = [d]
    _, state = run(repo, research_bundle)
    assert state["research"]["dividends"][0]["eligible_quantity"] == "40"
    assert Decimal(state["research"]["dividends"][0]["gross_entitlement"]) == 100


def test_net_return_uses_declared_tax_and_costs(repo, research_bundle):
    m = research_bundle["research"]["models"][0]
    m.update(dividend_tax_rate="0.2", capital_gains_tax_rate="0.1", entry_cost_rate="0.01",
             exit_cost_rate="0.01", liabilities_per_share="1")
    for s in m["scenarios"]:
        s["exit_price"] = "150"
    _, state = run(repo, research_bundle)
    row = next(r for r in state["research"]["valuations"] if r["ticker"] == "DEMO-BANK")
    # Paid 126.25; proceeds 148.5 - 2.5 + 4 - 1 = 149.
    expected = (Decimal(149) - Decimal("126.25")) / Decimal("126.25")
    assert abs(Decimal(row["net_return"]) - expected) < Decimal("1e-25")


def test_successful_research_packet_import(repo, research_bundle):
    research = research_bundle.pop("research")
    sources = [research_bundle["sources"].pop()]
    run(repo, research_bundle)
    result = pipeline.import_research(repo, dumps({"sources": sources, "research": research}))
    assert result["status"] == "PARTIAL"
    assert read_json(repo, pipeline.STATE)["research"]["summary"]["eligible_count"] == 3


def test_tampered_phase2_workbook_is_rejected(repo, research_bundle, tmp_path):
    from investment_workspace.excel import validate_workbook
    _, state = run(repo, research_bundle)
    wb = load_workbook(repo / pipeline.WORKBOOK)
    ws = wb["Valuation"]
    columns = [c.value for c in ws[1]]
    ws.cell(2, columns.index("present_fair_value") + 1, 1)
    path = tmp_path / "tampered.xlsx"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="parity"):
        validate_workbook(path, state)


def test_universe_revision_keeps_prior_source_binding(repo, research_bundle):
    _, before = run(repo, research_bundle)
    source = deepcopy(research_bundle["sources"][-1])
    source["source_id"] = "synthetic.universe.2"
    research_bundle["sources"].append(source)
    research_bundle["research"]["universe_source_ids"] = [source["source_id"]]
    research_bundle["research"]["required_tickers"].append("MISSING-FILING")
    result, after = run(repo, research_bundle)
    assert result["status"] == "PARTIAL"
    assert any(c["ticker"] == "MISSING-FILING" and c["outcome"] == "BLOCKED_MODEL" for c in after["coverage"])
    assert len(after["research"]["history"]) == len(before["research"]["history"])


def test_recovered_model_must_reproduce_historical_outputs(repo, research_bundle):
    m = research_bundle["research"]["models"][0]
    m.update(origin="LEGACY_RECOVERED", legacy_column_map={"old.bvps": "book_value_per_share"},
             historical_outputs={"bear": "162.5", "base": "225", "bull": "287.5"})
    _, state = run(repo, research_bundle)
    assert any(r["ticker"] == "DEMO-BANK" for r in state["research"]["valuations"])


def test_valuation_without_research_is_partial(repo, bundle):
    pipeline.run(repo, dumps(bundle))
    result = pipeline.run(repo, pipeline.current_input(repo), workflow="valuation")
    assert result["status"] == "PARTIAL"
    assert main(["--root", str(repo), "run", "valuation", "--require-full-coverage"]) == 1


def test_phase2_init_in_empty_directory(tmp_path):
    root = tmp_path / "new-demo"
    assert main(["--root", str(root), "init", "--offline", "--phase2"]) == 0
    assert read_json(root, pipeline.STATE)["research"]["summary"]["eligible_count"] == 3
    assert not storage.git(root, "status", "--porcelain").stdout
