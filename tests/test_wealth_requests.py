"""Chat requests exercise the actual calculator and journal without external writes."""

import json
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from investment_workspace import pipeline, storage
from investment_workspace.storage import dumps, git, read_json, head
from investment_workspace.wealth_requests import execute, pending, preflight, gate
from investment_workspace.wealth import compare_histories


@pytest.fixture
def ready(repo):
    pipeline.run(repo,pipeline.fixture_bytes(phase4=True))
    return repo


def request(root, operation='simulate', name='test-1', package=None, **fields):
    state=read_json(root,pipeline.STATE)
    value={'request_id':name,'operation':operation,'dataset_id':state['dataset']['dataset_id'],
           'expected_run_id':state['published_run_id'],'as_of':'2026-10-09',
           'instruction':'Save these inputs' if operation=='accept-and-run' else 'Calculate a what-if scenario',
           'authorized_by':'synthetic test user',**fields}
    if package is not None: value['package']=package
    path=root/'requests'/'wealth'/(name+'.json')
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_bytes(dumps(value))
    git(root,'add','--',str(path.relative_to(root)))
    git(root,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','synthetic request')
    return path


def revised(root):
    data=json.loads(pipeline.current_input(root))
    wealth=deepcopy(data['wealth'])
    r=wealth['scenarios'][0]
    r['supersedes']=r['record_id'];r['record_id']+='-revised'
    r['returns'][0]['annual_nominal_return']='0.20'
    source=deepcopy(data['sources'][-1]);source['source_id']='synthetic.wealth.request-policy'
    r['source_ids']=[source['source_id']]
    return {'sources':[source],'wealth':wealth}


def test_simulation_isolated_and_idempotent(ready):
    path=request(ready,package=revised(ready))
    before=storage.committed(ready,pipeline.STATE)
    result=execute(ready,path,local_demo=True)
    assert result['status']=='COMPLETE'
    assert result['modules']=={'retirement':'COMPLETE','cross-asset':'COMPLETE'}
    assert result['accepted_inputs_changed'] is False
    assert storage.committed(ready,pipeline.STATE)==before
    assert (ready/result['outputs']['workbook']).is_file()
    commit=head(ready)
    assert execute(ready,path,local_demo=True)['idempotent'] is True
    assert head(ready)==commit
    assert pipeline.validate(ready)['integrity']=='PASS'


def test_accept_publishes_input_receipt_and_workbook_together(ready):
    packet=revised(ready)
    path=request(ready,'accept-and-run',package=packet)
    result=execute(ready,path,local_demo=True)
    assert result['accepted_inputs_changed'] is True
    accepted=json.loads(pipeline.current_input(ready))
    assert accepted['wealth']['scenarios'][0]['record_id']==packet['wealth']['scenarios'][0]['record_id']
    assert pipeline.validate(ready)['workbook_parity']=='PASS'
    names=git(ready,'diff-tree','--no-commit-id','--name-only','-r','HEAD').stdout.decode()
    assert 'reports/requests/test-1.json' in names
    assert 'data/canonical/portfolio.json' in names
    assert 'workbooks/Investment_Workspace.xlsx' in names


@pytest.mark.parametrize('fields,message',[
    ({'expected_run_id':'old-run'},'stale accepted revision'),
    ({'dataset_id':'other-dataset'},'dataset does not match'),
    ({'operation':'execute-shell','command':'echo unsafe'},'literal_error'),
    ({'modules':['retirement','retirement']},'duplicate requested module'),
])
def test_invalid_request_receipts_preserve_last_good(ready,fields,message):
    before=storage.committed(ready,pipeline.STATE)
    path=request(ready,**fields)
    result=execute(ready,path,local_demo=True)
    assert result['status']=='INVALID'
    assert message in str(result['errors'])
    assert storage.committed(ready,pipeline.STATE)==before


def test_mutated_request_is_rejected(ready):
    path=request(ready)
    execute(ready,path,local_demo=True)
    value=json.loads(path.read_bytes());value['instruction']='Different request'
    path.write_bytes(dumps(value));git(ready,'add','--',str(path.relative_to(ready)))
    git(ready,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','tamper request')
    with pytest.raises(ValueError,match='immutable'):
        execute(ready,path,local_demo=True)


def test_pending_scans_all_requests_and_does_not_repeat(ready):
    request(ready,name='a');request(ready,name='b')
    results=pending(ready,local_demo=True)
    assert len(results['requests'])==2
    commit=head(ready)
    assert all(r['idempotent'] for r in pending(ready,local_demo=True)['requests'])
    assert head(ready)==commit


def test_preflight_is_read_only_and_reports_missing_assumptions(ready):
    packet=revised(ready)
    packet['wealth']['scenarios'][0]['returns'][0]['return_tax_rate']=None
    commit=head(ready)
    result=preflight(ready,packet)
    assert result['status']=='PARTIAL'
    assert any(g['reason']=='MISSING_RETIREMENT_INPUTS' for g in result['missing_inputs'])
    assert head(ready)==commit
    assert not git(ready,'status','--porcelain').stdout


def test_validate_request_does_not_export_or_accept(ready):
    before=storage.committed(ready,pipeline.STATE)
    result=execute(ready,request(ready,'validate'),local_demo=True)
    assert result['status']=='COMPLETE'
    assert 'outputs' not in result
    assert storage.committed(ready,pipeline.STATE)==before


def test_private_policy_and_runner_required_before_read(ready,monkeypatch):
    path=request(ready)
    before=head(ready)
    monkeypatch.setenv('GITHUB_ACTIONS','true')
    monkeypatch.setenv('WEALTH_PRIVATE_REPOSITORY','false')
    with pytest.raises(ValueError,match='privacy'):
        execute(ready,path)
    assert head(ready)==before
    with pytest.raises(ValueError,match='authorization'):
        pending(ready,local_demo=True,push=True)


def test_atomic_request_commit_failure_rolls_back_and_retries(ready,monkeypatch):
    path=request(ready,'accept-and-run',package=revised(ready))
    before=storage.committed(ready,pipeline.STATE)
    original=storage.git
    def fail(root,*args,**kwargs):
        if root==ready and 'commit' in args:
            raise OSError('injected request commit failure')
        return original(root,*args,**kwargs)
    monkeypatch.setattr(storage,'git',fail)
    with pytest.raises(storage.PublicationError,match='interrupted'):
        execute(ready,path,local_demo=True)
    assert storage.committed(ready,pipeline.STATE)==before
    monkeypatch.setattr(storage,'git',original)
    result=execute(ready,path,local_demo=True)
    assert result['status']=='COMPLETE'
    assert not (ready/'.runtime/pending.json').exists()
    assert pipeline.validate(ready)['workbook_parity']=='PASS'


def test_common_period_comparison_never_uses_different_endpoints():
    from types import SimpleNamespace
    lh=rh=SimpleNamespace(return_basis='PRICE')
    left={date(2026,1,1):Decimal(10),date(2026,2,1):Decimal(20),date(2026,3,1):Decimal(40),date(2026,4,1):Decimal(80)}
    right={date(2026,2,1):Decimal(100),date(2026,3,1):Decimal(110),date(2026,4,1):Decimal(120),date(2026,5,1):Decimal(400)}
    result=compare_histories(lh,left,rh,right,3)
    assert result['common_start']=='2026-02-01'
    assert result['common_end']=='2026-04-01'
    assert result['left_common_return']==3
    assert result['right_common_return']==Decimal('.2')
    assert result['status']=='INSUFFICIENT_HISTORY_OR_BASIS_MISMATCH'


def test_cross_asset_budget_resume_and_history_invalidation(ready):
    raw=pipeline.current_input(ready)
    # Explicit full verification spends budget on retirement first, leaving pair deferred.
    pipeline.run(ready,raw,max_units=0,recalculate=True)
    state=read_json(ready,pipeline.STATE)
    assert any(r['owner']=='cross-asset' for r in state['deferred_work'])
    assert state['wealth']['summary']['cross_asset_status']=='PARTIAL'
    pipeline.run(ready,raw)
    state=read_json(ready,pipeline.STATE)
    assert state['wealth']['summary']['cross_asset_status']=='COMPLETE'
    assert not state['deferred_work']


@pytest.fixture
def private_runner(ready,tmp_path,monkeypatch):
    policy={'enabled':True,'repository':'test/wealth','authorized_by':'owner',
            'auto_push':True,'max_work_units':2000,'max_seconds':300,'max_request_bytes':5242880}
    path=ready/'config/wealth-requests.json';path.parent.mkdir(exist_ok=True);path.write_bytes(dumps(policy))
    git(ready,'add','--','config/wealth-requests.json')
    git(ready,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','enable synthetic test policy')
    remote=tmp_path/'remote.git';git(ready,'init','--bare','-q',str(remote))
    git(ready,'remote','add','origin',str(remote))
    git(ready,'push','-q','origin','HEAD')
    monkeypatch.setenv('GITHUB_ACTIONS','true')
    monkeypatch.setenv('GITHUB_REPOSITORY','test/wealth')
    monkeypatch.setenv('WEALTH_PRIVATE_REPOSITORY','true')
    return ready,remote


def publish_request(root):
    path=request(root)
    git(root,'push','-q','origin','HEAD')
    return path


def test_private_runner_publishes_and_verifies_remote(private_runner):
    root,remote=private_runner
    publish_request(root)
    result=pending(root,push=True)
    assert result['status']=='PROCESSED'
    assert result['publication']['status']=='PUSHED'
    assert head(root)==git(remote,'rev-parse','HEAD').stdout.decode().strip()
    assert read_json(root,'reports/requests/test-1.json')['status']=='COMPLETE'


def test_remote_divergence_blocks_before_request_calculation(private_runner):
    root,remote=private_runner
    request(root)
    before=head(root)
    with pytest.raises(storage.PublicationError,match='remote HEAD differs'):
        pending(root,push=True)
    assert head(root)==before
    assert read_json(root,'reports/requests/test-1.json') is None


def test_normal_push_failure_retains_committed_receipt(private_runner,monkeypatch):
    root,remote=private_runner
    publish_request(root)
    def fail(*args,**kwargs):
        raise storage.PublicationError('injected push failure; local commit retained')
    monkeypatch.setattr(pipeline,'push',fail)
    with pytest.raises(storage.PublicationError,match='push failure'):
        pending(root,push=True)
    assert read_json(root,'reports/requests/test-1.json')['status']=='COMPLETE'
    assert pipeline.validate(root)['workbook_parity']=='PASS'
    assert head(root)!=git(remote,'rev-parse','HEAD').stdout.decode().strip()


def test_malformed_request_does_not_starve_later_requests(ready):
    path=request(ready,name='a')
    path.write_text('{')
    git(ready,'add','--',str(path.relative_to(ready)))
    git(ready,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','malformed synthetic request')
    request(ready,name='b')
    result=pending(ready,local_demo=True)
    assert [r['status'] for r in result['requests']]==['INVALID','COMPLETE']
    assert result['status']=='PARTIAL'


def test_retirement_timeout_is_recorded_as_deferred(ready,monkeypatch):
    from investment_workspace.runtime import RunContext,context
    from investment_workspace.schemas import Bundle
    from investment_workspace.wealth import calculate_wealth
    bundle=Bundle.model_validate_json(pipeline.current_input(ready))
    previous=read_json(ready,pipeline.STATE)
    state=deepcopy(previous);state['data_gaps']=[]
    runtime=RunContext(ready,pipeline.code_hash(),{},max_seconds=0.01)
    original=runtime.consume
    def consume(owner,key):
        original(owner,key)
        if owner=='retirement': runtime.began-=1
    monkeypatch.setattr(runtime,'consume',consume)
    with context(runtime):
        calculate_wealth(bundle,state,previous,'timeout-test','hash')
    assert any(d['owner']=='retirement' for d in runtime.deferred)
    assert state['wealth']['summary']['retirement_status']=='PARTIAL'


def test_benchmark_parity_and_no_caller_mutation(ready):
    from investment_workspace.wealth_benchmark import benchmark
    before=head(ready)
    result=benchmark(samples=1,months=2)
    assert result['parity']=='PASS' and result['ai_cost'] is None
    assert len(result['measurements'])==4
    assert result['unchanged']['full_median_seconds']>0
    assert head(ready)==before
    with pytest.raises(ValueError): benchmark(samples=0)


def test_changed_history_recalculates_only_its_pair(ready):
    data=json.loads(pipeline.current_input(ready))
    before=read_json(ready,pipeline.STATE)
    source=deepcopy(data['sources'][-1]);source['source_id']='synthetic.wealth.history-update'
    data['sources'].append(source)
    history=data['wealth']['histories'][1]
    history.update(supersedes=history['record_id'],record_id=history['record_id']+'-updated',source_ids=[source['source_id']])
    history['quotes'][-1]['close']=str(Decimal(history['quotes'][-1]['close'])+1)
    pipeline.run(ready,dumps(data))
    after=read_json(ready,pipeline.STATE)
    assert after['usage'][-1]['calculation_units']==1
    assert after['wealth']['projections']==before['wealth']['projections']
    assert after['wealth']['correlations']!=before['wealth']['correlations']


def test_module_status_is_independent_of_missing_other_module(ready):
    data=json.loads(pipeline.current_input(ready))
    # A new candidate dataset omits history only before its first import.
    from investment_workspace.wealth_requests import summary
    state=read_json(ready,pipeline.STATE)
    state['wealth']['summary']['cross_asset_status']='PARTIAL'
    assert summary(state,['retirement'])['status']=='COMPLETE'
    assert summary(state,['retirement','cross-asset'])['status']=='PARTIAL'


def test_request_changed_before_first_processing_cannot_be_accepted(ready):
    path=request(ready)
    raw=json.loads(path.read_bytes());raw['operation']='accept-and-run'
    path.write_bytes(dumps(raw));git(ready,'add','--',str(path.relative_to(ready)))
    git(ready,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','edit submitted request')
    before=storage.committed(ready,pipeline.STATE)
    result=execute(ready,path,local_demo=True)
    assert result['status']=='INVALID'
    assert 'immutable' in str(result['errors'])
    assert result['request_hash']!=result['submitted_hash']
    assert storage.committed(ready,pipeline.STATE)==before


def test_wealth_dashboard_is_checked_and_summary_amounts_are_numeric(ready,tmp_path):
    from openpyxl import load_workbook
    from investment_workspace.excel import validate_workbook
    state=read_json(ready,pipeline.STATE)
    wb=load_workbook(ready/'workbooks/Investment_Workspace.xlsx')
    sheet=wb['Retirement_Summary'];headers=[c.value for c in sheet[1]]
    assert isinstance(sheet.cell(2,headers.index('starting_wealth_pkr')+1).value,(int,float))
    assert isinstance(sheet.cell(2,headers.index('ending_accessible_balance')+1).value,(int,float))
    for row in wb['Dashboard'].iter_rows(min_row=2):
        if row[0].value=='wealth_retirement_status': row[1].value='COMPLETE_BUT_TAMPERED'
    path=tmp_path/'tampered-dashboard.xlsx';wb.save(path);wb.close()
    with pytest.raises(ValueError,match='wealth dashboard parity'):
        validate_workbook(path,state)
