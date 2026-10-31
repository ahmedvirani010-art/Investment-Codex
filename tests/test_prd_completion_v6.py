"""Audit regressions, shared contracts and independent numerical checks."""
import io
import json
from copy import deepcopy
from datetime import date,datetime,timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from investment_workspace import pipeline,storage
from investment_workspace.schemas import Bundle,ScenarioOverride
from investment_workspace.portfolio import calculate
from investment_workspace.research import calculate_research,dividend_result
from investment_workspace.runtime import RunContext,context
from investment_workspace.scenarios import resolve
from investment_workspace.sector_engines import forecast
from investment_workspace.models import reference_value
from investment_workspace.decisions import basis,calculate as decisions
from investment_workspace.object_storage import StorageRequest,PullObject,execute as transfer_execute,transfer,artifact
from investment_workspace.workspace_requests import execute,pending
from investment_workspace.storage import dumps,read_json,git,head,digest


@pytest.fixture
def full():
    return json.loads(Path('src/investment_workspace/fixtures/completion-demo.json').read_bytes())


def publish(repo,full):
    result=pipeline.run(repo,dumps(full))
    assert result['status']!='FAILED',result.get('error')
    return read_json(repo,pipeline.STATE)


def test_dividend_roll_preserves_verified_entitlement_and_cash(full):
    b=Bundle.model_validate(full);state=calculate(b);d=b.research.dividends[0]
    first=dividend_result(d,b,state,{'run_id':'first'},lambda _:True)
    assert first['status']=='RECEIVED'
    rolled=deepcopy(full);rolled['market_cutoff']='2026-10-20'
    rolled['opening']=deepcopy(rolled['closing']);rolled['opening']['timestamp']='2026-10-13T18:00:00+05:00'
    rolled['closing']['timestamp']='2026-10-20T18:00:00+05:00';rolled['transactions']=[]
    next_bundle=Bundle.model_validate(rolled)
    second=dividend_result(d,next_bundle,state,{'run_id':'next'},lambda _:True,{'dividends':[first]})
    assert second['status']=='RECEIVED'
    assert second['eligible_quantity']==first['eligible_quantity']
    assert second['entitlement_basis']==first['entitlement_basis']
    assert Decimal(second['net_pending'])==0
    assert state['summary']['cash']=='15220.00'


def test_partial_receipts_and_unknown_tax_do_not_invent_pending(full):
    b=Bundle.model_validate(full);state=calculate(b);d=b.research.dividends[0]
    tx=next(t for t in state['transactions'] if t['transaction_id'] in d.receipt_transaction_ids)
    tx['amount']='100';tx['taxes']='20';tx['fees']='3'
    r=dividend_result(d,b,state,{},lambda _:True)
    assert r['status']=='PARTIALLY_RECEIVED'
    assert Decimal(r['net_pending'])==120
    assert Decimal(r['net_received'])==77
    d.tax_rate=None
    r=dividend_result(d,b,state,{},lambda _:True)
    assert r['net_pending'] is None
    assert r['tax_status']=='UNKNOWN'


def test_correction_outside_window_keeps_previous_proof(full):
    b=Bundle.model_validate(full);state=calculate(b);d=b.research.dividends[0]
    first=dividend_result(d,b,state,{},lambda _:True)
    revised=d.model_copy(update={'record_id':'corrected','supersedes':d.record_id,'dps':Decimal('9')})
    moved=b.model_copy(update={'opening':b.opening.model_copy(update={'timestamp':datetime(2026,10,10,tzinfo=timezone.utc)})})
    out=dividend_result(revised,moved,state,{},lambda _:True,{'dividends':[first]})
    assert out['status']=='CORRECTION_PENDING'
    assert out['previous_verified_result']['entitlement_basis']==first['entitlement_basis']


@pytest.mark.parametrize('idx,pat,eps,fcfe',[(0,'50','5',None),(1,'100','10','75'),(2,'70','7','54')])
def test_first_wave_engines_hand_calculated(full,idx,pat,eps,fcfe):
    b=Bundle.model_validate(full);m=b.research.models[idx];f=m.scenarios[0].sector_forecasts[0]
    row=forecast(f,m,b)
    assert row['pat']==Decimal(pat) and row['eps']==Decimal(eps)
    if fcfe: assert row['fcfe']==Decimal(fcfe)
    if idx==0:
        assert row['closing_equity']==540
        assert row['dividend_capacity']==80
        assert row['oci']==10
    assert reference_value(m,m.scenarios[0],f.engine,b)>0


@pytest.mark.parametrize('fault', ['capital','reserves','collections','unit','approval','cash'])
def test_numerical_engine_guardrails(full,fault):
    b=Bundle.model_validate(full)
    if fault=='capital':
        m=b.research.models[0];m.scenarios[0].sector_forecasts[0].proposed_dividends=Decimal(1000)
    else:
        m=b.research.models[1]
        f=m.scenarios[0].sector_forecasts[0]
        if fault=='reserves': f.commodities[0].production=Decimal(1001)
        if fault=='collections': f.collections=Decimal(10000)
        if fault=='unit': f.commodities[0].fx_to_pkr=Decimal(2)
        if fault=='approval': m.approval=None
        if fault=='cash': f.cash_income_included=True;m.scenarios[0].surplus_cash_per_share=Decimal(2);m.scenarios[0].cash_income_excluded=True
    with pytest.raises(ValueError): reference_value(m,m.scenarios[0],m.scenarios[0].sector_forecasts[0].engine,b)


def test_shared_profile_overrides_units_expiry_and_dependency(full):
    b=Bundle.model_validate(full);s=b.research.models[0].scenarios[0]
    assert resolve(b,s)['status']=='LINKED'
    s.discount_rate=Decimal('.15')
    with pytest.raises(ValueError,match='override'): resolve(b,s)
    s.profile.overrides['discount_rate']=ScenarioOverride.model_validate({'assumption':{'value':'.15','unit':'FRACTION'},'rationale':'explicit override','source_ids':s.sector_forecasts[0].source_ids})
    assert resolve(b,s)['effective_assumptions']['discount_rate']['value']=='0.15'
    s.profile.overrides['discount_rate'].assumption.unit='PKR'
    with pytest.raises(ValueError,match='units'): resolve(b,s)
    b=Bundle.model_validate(full);b.scenario_profiles[0].valid_until=date(2026,10,8)
    with pytest.raises(ValueError,match='stale'): resolve(b,b.research.models[0].scenarios[0])


def test_numerical_fixture_pipeline_and_uniform_lineage(repo,full):
    state=publish(repo,full)
    assert state['research']['summary']['valuation_status']=='COMPLETE'
    assert len(state['research']['forecasts'])==9
    assert all(r['forecast_contract']=='NUMERICAL_V1' for r in state['research']['valuations'][:9])
    assert all(r.get('code_version') and r.get('code_hash') for module in ('research','integration','wealth') for r in state[module][{'research':'valuations','integration':'timing','wealth':'projections'}[module]])
    assert pipeline.validate(repo)['workbook_parity']=='PASS'
    publish(repo,full)
    assert read_json(repo,pipeline.STATE)['attempt_status']=='NO_MATERIAL_CHANGE'


@pytest.mark.parametrize('offset,expected',[('2026-10-09',11),('2026-10-10',10),('2026-10-20',0)])
def test_catalyst_uses_assessment_not_old_cutoff(repo,full,offset,expected):
    result=pipeline.run(repo,dumps(full),as_of=offset)
    assert result['status']!='FAILED'
    state=read_json(repo,pipeline.STATE)
    assert state['integration']['catalysts'][0]['days_to_start']==expected
    assert state['summary']['market_cutoff']=='2026-10-09'


def test_decisions_approval_link_and_changed_basis(repo,full):
    state=publish(repo,full);b=Bundle.model_validate(full)
    proposal=b.integration.proposals[0]
    record={'record_id':'decision.1','ticker':proposal.ticker,'effective_date':'2026-10-09','valid_until':'2026-10-12',
            'source_ids':proposal.source_ids,'proposal_id':proposal.record_id,'decision':'APPROVE','decided_by':'synthetic owner',
            'decided_at':'2026-10-09T00:00:00+05:00','reason':'explicit approval','reviewed_basis_hash':basis(state,proposal.ticker),
            'approved_quantity':'10','transaction_ids':[]}
    from investment_workspace.schemas import ICDecision
    b.decisions=[ICDecision.model_validate(record)]
    out=deepcopy(state);decisions(b,out,{},'decision-run','hash')
    assert out['decisions']['rows'][0]['review_status']=='CURRENT'
    changed=deepcopy(out);changed['holdings'][0]['price']='999'
    decisions(b,changed,out,'decision-next','hash')
    assert changed['decisions']['rows'][0]['review_status']=='NEEDS_REVIEW'
    assert changed['decisions']['rows'][0]['orders_created'] is False


class Missing(Exception):
    response={'Error':{'Code':'NoSuchKey'}}


class FakeS3:
    def __init__(self): self.data={};self.puts=0;self.fail=0
    def get_object(self,**kw):
        if self.fail: self.fail-=1;raise OSError('temporary service failure')
        if kw['Key'] not in self.data: raise Missing()
        value=self.data[kw['Key']]
        return {'Body':io.BytesIO(value),'ContentLength':len(value)}
    def head_object(self,**kw):
        if kw['Key'] not in self.data: raise Missing()
        return {}
    def put_object(self,**kw): self.puts+=1;self.data[kw['Key']]=kw['Body'];return {}


@pytest.fixture
def storage_ready(repo,full,monkeypatch):
    publish(repo,full)
    p=json.loads(Path('config/storage.json').read_text());p['enabled']=True
    (repo/'config').mkdir(exist_ok=True);(repo/'config/storage.json').write_bytes(dumps(p))
    git(repo,'add','--','config/storage.json');git(repo,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','synthetic transfer policy')
    monkeypatch.setenv('R2_BUCKET','synthetic-bucket')
    return repo,FakeS3()


def test_r2_pull_pins_artifact_and_does_not_accept_financial_records(storage_ready):
    repo,s3=storage_ready;data=b'{"synthetic":"only"}';s3.data['approved/data.json']=data
    before=storage.committed(repo,pipeline.STATE)
    req=StorageRequest(request_id='pull.1',operation='pull',objects=[PullObject(key='approved/data.json',sha256=digest(data))])
    r=transfer_execute(repo,req,s3=s3)
    assert r['status']=='COMPLETE'
    assert artifact(repo,'pull.1',digest(data))==data
    assert storage.committed(repo,pipeline.STATE)==before
    assert transfer_execute(repo,req,s3=s3)['idempotent']


@pytest.mark.parametrize('key',['../secret','approved/../secret','https://evil/file','/approved/file','other/file'])
def test_r2_rejects_paths_and_prefixes(storage_ready,key):
    repo,s3=storage_ready
    with pytest.raises(ValueError): transfer(repo,StorageRequest(request_id='bad',operation='pull',objects=[PullObject(key=key)]),s3=s3)


def test_r2_hash_mismatch_unpinned_and_verified_archives(storage_ready):
    repo,s3=storage_ready;s3.data['sources/file']=b'actual'
    wrong=StorageRequest(request_id='wrong',operation='pull',objects=[PullObject(key='sources/file',sha256='0'*64)])
    assert transfer_execute(repo,wrong,s3=s3)['status']=='FAILED'
    r=transfer_execute(repo,StorageRequest(request_id='unpinned',operation='pull',objects=[PullObject(key='sources/file')]),s3=s3)
    with pytest.raises(ValueError,match='pinned'): artifact(repo,'unpinned',r['objects'][0]['artifact_id'])
    run=read_json(repo,pipeline.STATE)['published_run_id']
    req=StorageRequest(request_id='archive',operation='push',run_id=run,paths=['workbooks/Investment_Workspace.xlsx'])
    result=transfer_execute(repo,req,s3=s3)
    assert result['status']=='COMPLETE'
    assert result['archive_manifest']['key'] in s3.data
    assert result['financial_acceptance'] is False
    puts=s3.puts;transfer_execute(repo,req,s3=s3);assert s3.puts==puts


def request(repo,operation,**kw):
    state=read_json(repo,pipeline.STATE);rid=kw.pop('request_id','workspace.1')
    spec={'request_id':rid,'operation':operation,'dataset_id':state['dataset']['dataset_id'],'expected_run_id':state['published_run_id'],
          'as_of':'2026-10-09','instruction':'synthetic instruction','authorized_by':'synthetic test owner',**kw}
    path=repo/'requests/workspace'/(rid+'.json');path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(dumps(spec))
    git(repo,'add','--',str(path.relative_to(repo)));git(repo,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','synthetic workspace request')
    return path


@pytest.mark.parametrize('workflow',['portfolio','valuation','update-dividends','timing','macro','watchlist','ic','retirement','cross-asset','daily','event'])
def test_every_workspace_module_request_is_isolated(repo,full,workflow):
    publish(repo,full);before=storage.committed(repo,pipeline.STATE)
    r=execute(repo,request(repo,'simulate',workflows=[workflow]),local_demo=True)
    assert r['status'] in {'COMPLETE','PARTIAL'},r['errors']
    assert workflow in r['modules']
    assert storage.committed(repo,pipeline.STATE)==before
    assert r['accepted_inputs_changed'] is False


def test_workspace_r2_request_and_explicit_import(storage_ready):
    repo,s3=storage_ready
    state=read_json(repo,pipeline.STATE)
    packet={'sources':[],'decisions':[]}
    data=dumps(packet);s3.data['approved/package.json']=data
    path=request(repo,'storage.pull',storage={'request_id':'workspace.1','operation':'pull','objects':[{'key':'approved/package.json','sha256':digest(data)}]})
    result=execute(repo,path,local_demo=True,s3=s3)
    assert result['status']=='COMPLETE',result['errors']
    r=execute(repo,request(repo,'import',request_id='import.1',package_artifact={'receipt_id':'workspace.1','artifact_id':digest(data)},workflows=['valuation']),local_demo=True,s3=s3)
    assert r['status']=='COMPLETE',r['errors']
    assert pipeline.validate(repo)['integrity']=='PASS'


def test_workspace_arbitrary_path_code_and_public_execution_blocked(repo,full):
    publish(repo,full)
    path=request(repo,'simulate',package={'files':[{'path':'/etc/passwd'}]})
    r=execute(repo,path,local_demo=True)
    assert r['status']=='INVALID'
    with pytest.raises(ValueError,match='privacy'): execute(repo,path)


def test_workspace_archive_failure_keeps_valid_run(repo,full):
    publish(repo,full)
    path=request(repo,'run',workflows=['valuation'],archive_paths=['workbooks/Investment_Workspace.xlsx'])
    r=execute(repo,path,local_demo=True,s3=FakeS3())
    assert r['status']=='PARTIAL'
    assert r['archive']['status']=='FAILED'
    assert r['modules']['valuation']=='COMPLETE'
    assert pipeline.validate(repo)['integrity']=='PASS'


def test_sdk_configuration_is_sigv4_auto_without_network(monkeypatch):
    pytest.importorskip('boto3')
    from investment_workspace.object_storage import client
    monkeypatch.setenv('R2_ENDPOINT_URL','https://00000000000000000000000000000000.r2.cloudflarestorage.com')
    monkeypatch.setenv('R2_BUCKET','synthetic-bucket')
    monkeypatch.setenv('R2_ACCESS_KEY_ID','synthetic-access')
    monkeypatch.setenv('R2_SECRET_ACCESS_KEY','synthetic-secret')
    s3=client()
    assert s3.meta.config.region_name=='auto'
    assert s3.meta.config.signature_version=='s3v4'
    assert s3.meta.config.read_timeout==30
    s3.close()


def test_missing_numerical_driver_blocks_certification(full):
    full['research']['models'][0]['scenarios'][0]['sector_forecasts'][0]['tax_charge']=None
    b=Bundle.model_validate(full);m=b.research.models[0]
    with pytest.raises(ValueError,match='missing numerical'): reference_value(m,m.scenarios[0],'Bank',b)


def test_full_decision_transaction_lifecycle_and_no_double_links(repo,full):
    from investment_workspace.schemas import ICDecision
    state=publish(repo,full);b=Bundle.model_validate(full);p=b.integration.proposals[0]
    d={'record_id':'approval.exec','ticker':p.ticker,'effective_date':'2026-10-09','valid_until':'2026-10-12','source_ids':p.source_ids,
       'proposal_id':p.record_id,'decision':'APPROVE','decided_by':'synthetic owner','decided_at':'2026-10-09T00:00:00+05:00',
       'reason':'explicit review','reviewed_basis_hash':basis(state,p.ticker),'approved_quantity':'10','transaction_ids':['execution.1']}
    state['transactions'].append({'transaction_id':'execution.1','ticker':p.ticker,'kind':'BUY','quantity':'4','timestamp':'2026-10-09T12:00:00+05:00'})
    b.decisions=[ICDecision.model_validate(d)];decisions(b,state,{},'partial','hash')
    assert state['decisions']['rows'][0]['execution_status']=='PARTIALLY_EXECUTED'
    state['transactions'][-1]['quantity']='10';decisions(b,state,{},'full','hash')
    assert state['decisions']['rows'][0]['execution_status']=='EXECUTED'
    state['transactions'][-1]['quantity']='11';decisions(b,state,{},'excess','hash')
    assert state['decisions']['rows'][0]['execution_status']=='EXECUTION_EXCEPTION'
    d['decision']='REJECT'
    with pytest.raises(ValueError,match='only approvals'): ICDecision.model_validate(d)


def test_r2_retries_limits_interruption_and_corrupt_archive(storage_ready,monkeypatch):
    root,s3=storage_ready;s3.data['approved/bytes']=b'abcd';s3.fail=1
    monkeypatch.setattr('investment_workspace.object_storage.time.sleep',lambda _:None)
    r=transfer(root,StorageRequest(request_id='retry',operation='pull',objects=[PullObject(key='approved/bytes',sha256=digest(b'abcd'))]),s3=s3)
    assert r['status']=='COMPLETE'
    class Interrupted(FakeS3):
        def get_object(self,**kw):
            class Broken(io.BytesIO):
                def read(self,*args): raise OSError('interrupted stream')
            return {'Body':Broken(b'abcd'),'ContentLength':4}
    assert transfer_execute(root,StorageRequest(request_id='interrupted',operation='pull',objects=[PullObject(key='approved/bytes')]),s3=Interrupted())['status']=='FAILED'
    run=read_json(root,pipeline.STATE)['published_run_id']
    spec=StorageRequest(request_id='corrupt',operation='push',run_id=run,paths=['workbooks/Investment_Workspace.xlsx'])
    planned=transfer(root,spec,dry_run=True)['objects'][0]
    s3.data[planned['key']]=b'corrupt'
    with pytest.raises(ValueError,match='read-back'): transfer(root,spec,s3=s3)
    assert s3.data[planned['key']]==b'corrupt'
    with pytest.raises(ValueError,match='manifest'): transfer(root,StorageRequest(request_id='badpath',operation='push',run_id=run,paths=['/etc/passwd']),s3=s3)


def test_shared_assumption_revision_blocks_owned_forecast_not_mutates_eps(repo,full):
    before=publish(repo,full)
    source=deepcopy(full['sources'][-1]);source['source_id']='synthetic.changed-profile';full['sources'].append(source)
    p=full['scenario_profiles'][0];p.update(record_id='changed-profile',supersedes=p['record_id'],source_ids=[source['source_id']])
    p['assumptions']['discount_rate']['value']='0.2'
    after=publish(repo,full)
    assert after['research']['summary']['valuation_status']=='PARTIAL'
    assert after['research']['earnings']==before['research']['earnings']
    assert full['research']['models'][0]['scenarios'][0]['forward_eps']=='5'


def test_native_date_and_compound_status_formatting(repo,full):
    publish(repo,full)
    wb=load_workbook(repo/'workbooks/Investment_Workspace.xlsx')
    ws=wb['Holdings'];headers=[c.value for c in ws[1]]
    cell=ws.cell(2,headers.index('price_date')+1)
    assert isinstance(cell.value,datetime) and cell.number_format=='yyyy-mm-dd'
    assert wb['Retirement_Summary'].max_row==4
    wb.close()


def test_explicit_portfolio_import_and_simulation(repo,full):
    state=publish(repo,full)
    portfolio={k:deepcopy(full[k]) for k in ('securities','opening','closing','transactions','prices')}
    portfolio['policy']=deepcopy(full.get('policy',{}))
    new_source=deepcopy(full['sources'][0]);new_source['source_id']='synthetic.new-prices'
    for p in portfolio['prices']: p['source_id']=new_source['source_id']
    portfolio['prices'][0]['close']='126'
    packet={'sources':[new_source],'portfolio':portfolio}
    before=storage.committed(repo,pipeline.STATE)
    simulated=execute(repo,request(repo,'simulate',request_id='positions.whatif',package=packet,workflows=['portfolio']),local_demo=True)
    assert simulated['status'] in {'COMPLETE','PARTIAL'},simulated['errors']
    assert storage.committed(repo,pipeline.STATE)==before
    accepted=execute(repo,request(repo,'accept-and-run',request_id='positions.accept',package=packet,workflows=['portfolio']),local_demo=True)
    assert accepted['status'] in {'COMPLETE','PARTIAL'},accepted['errors']
    assert accepted['accepted_inputs_changed']
    assert json.loads(pipeline.current_input(repo))['prices'][0]['close']=='126'


def test_workspace_normal_push_and_divergence_preserve_results(repo,full,tmp_path,monkeypatch):
    publish(repo,full)
    policy={'enabled':True,'repository':'synthetic/workspace','authorized_by':'test owner','auto_push':True,'max_work_units':2000,'max_seconds':300,'max_request_bytes':5242880}
    (repo/'config').mkdir(exist_ok=True);(repo/'config/workspace-requests.json').write_bytes(dumps(policy))
    git(repo,'add','--','config/workspace-requests.json');git(repo,'-c','user.name=Test','-c','user.email=test@localhost','commit','-qm','private test policy')
    remote=tmp_path/'remote.git';git(repo,'init','--bare','-q',str(remote));git(repo,'remote','add','origin',str(remote))
    request(repo,'simulate',workflows=['valuation'])
    git(repo,'push','-q','origin','HEAD')
    monkeypatch.setenv('GITHUB_ACTIONS','true');monkeypatch.setenv('GITHUB_REPOSITORY','synthetic/workspace');monkeypatch.setenv('WORKSPACE_PRIVATE_REPOSITORY','true')
    r=pending(repo,push=True)
    assert r['publication']['status']=='PUSHED'
    assert head(repo)==git(remote,'rev-parse','HEAD').stdout.decode().strip()
    request(repo,'simulate',request_id='unpublished',workflows=['valuation'])
    with pytest.raises(storage.PublicationError,match='remote HEAD'): pending(repo,push=True)
    assert read_json(repo,'reports/requests/unpublished.json') is None


def test_duplicate_forecast_period_cannot_double_count_cashflow(full):
    b=Bundle.model_validate(full);m=b.research.models[1];s=m.scenarios[0]
    s.sector_forecasts.append(s.sector_forecasts[0].model_copy(deep=True));s.fcfe.append(s.fcfe[0])
    with pytest.raises(ValueError,match='duplicate sector forecast period'): reference_value(m,s,'E&P',b)


def test_numerical_ep_requires_enforced_crosscheck(full):
    b=Bundle.model_validate(full);m=b.research.models[1];m.scenarios[0].crosscheck_tolerance=None
    with pytest.raises(ValueError,match='enforced payout'): reference_value(m,m.scenarios[0],'E&P',b)


def test_bank_mixed_cashflow_and_missing_share_inputs_fail_closed(full):
    b=Bundle.model_validate(full);m=b.research.models[0];s=m.scenarios[0]
    s.fcfe=[Decimal('1')]
    with pytest.raises(ValueError,match='FCFE'): reference_value(m,s,'Bank',b)
    s.fcfe=[];s.shares=None
    with pytest.raises(ValueError,match='share count'): reference_value(m,s,'Bank',b)


def test_workspace_decision_import_exposes_review_status(repo,full):
    state=publish(repo,full);p=state['integration']['input']['proposals'][0]
    source=deepcopy(full['sources'][-1]);source['source_id']='synthetic.decision-source'
    d={'record_id':'decision.import','ticker':p['ticker'],'effective_date':'2026-10-09','valid_until':'2026-10-12','source_ids':[source['source_id']],
       'proposal_id':p['record_id'],'decision':'APPROVE','decided_by':'synthetic owner','decided_at':'2026-10-09T00:00:00+05:00',
       'reason':'explicit historical approval','reviewed_basis_hash':'0'*64}
    r=execute(repo,request(repo,'decision',package={'sources':[source],'decisions':[d]},workflows=['ic']),local_demo=True)
    assert r['status']=='PARTIAL',r['errors']
    assert r['modules']['ic_decisions']=='PARTIAL'
    assert r['decisions'][0]['review_status']=='NEEDS_REVIEW'
    assert r['accepted_inputs_changed']
    assert read_json(repo,pipeline.STATE)['decisions']['rows'][0]['decision']=='APPROVE'
    assert pipeline.validate(repo)['workbook_parity']=='PASS'


def test_pinned_artifact_rehydrates_without_reaccepting_financial_data(storage_ready,monkeypatch):
    root,s3=storage_ready;data=b'{"sources":[],"decisions":[]}'
    s3.data['approved/rehydrate.json']=data
    r=transfer_execute(root,StorageRequest(request_id='rehydrate.source',operation='pull',objects=[PullObject(key='approved/rehydrate.json',sha256=digest(data))]),s3=s3)
    path=root/'.runtime/r2/artifacts'/digest(data);path.unlink()
    monkeypatch.setattr('investment_workspace.object_storage.client',lambda:s3)
    before=storage.committed(root,pipeline.STATE)
    assert artifact(root,'rehydrate.source',digest(data))==data
    assert storage.committed(root,pipeline.STATE)==before
    path.unlink();s3.data['approved/rehydrate.json']=b'tampered'
    with pytest.raises(ValueError,match='SHA-256'): artifact(root,'rehydrate.source',digest(data))
