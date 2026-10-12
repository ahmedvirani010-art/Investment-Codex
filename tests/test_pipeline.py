import json
import socket
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from investment_workspace import pipeline, storage
from investment_workspace.cli import main
from investment_workspace.excel import NOTE_COLUMNS, validate_workbook
from investment_workspace.schemas import NoteOverride
from investment_workspace.storage import committed, dumps, git, read_json, writer_lock


def run(repo, bundle):
    return pipeline.run(repo, dumps(bundle))


def test_offline_workbook_lineage_and_hashes(repo, bundle):
    result = run(repo, bundle)
    assert result["status"] == "SUCCESS"
    assert result["publication_status"] == "LOCAL_ONLY"
    state = read_json(repo, pipeline.STATE)
    assert state["published_run_id"] == result["run_id"]
    assert len(state["holdings"]) == 4
    for holding in state["holdings"]:
        assert holding["source_ids"]
        assert holding["input_hash"] == result["input_hash"]
        assert holding["run_id"] == result["run_id"]
    assert pipeline.validate(repo)["workbook_parity"] == "PASS"
    assert not git(repo, "status", "--porcelain").stdout
    assert git(repo, "ls-files", pipeline.WORKBOOK).stdout
    wb = load_workbook(repo / pipeline.WORKBOOK)
    assert "SYNTHETIC" in wb["Dashboard"]["B2"].value
    assert len(wb.sheetnames) >= 25
    wb.close()


def test_no_change_retains_dates_records_and_does_not_duplicate_ledger(repo, bundle):
    first = run(repo, bundle)
    initial = read_json(repo, pipeline.STATE)
    bundle["transactions"].append(deepcopy(bundle["transactions"][1]))
    second = run(repo, bundle)
    state = read_json(repo, pipeline.STATE)
    assert second["status"] == "NO_MATERIAL_CHANGE"
    assert state["holdings"] == initial["holdings"]
    assert state["analysis_run_id"] == first["run_id"]
    assert len(state["transactions"]) == 5
    assert len(state["run_log"]) == 2
    assert pipeline.validate(repo, first["run_id"])["integrity"] == "PASS"


def test_failed_attempt_updates_log_preserves_last_good(repo, bundle):
    first = run(repo, bundle)
    initial = read_json(repo, pipeline.STATE)
    bundle["transactions"][2]["quantity"] = "999"
    failed = run(repo, bundle)
    state = read_json(repo, pipeline.STATE)
    assert failed["status"] == "FAILED"
    assert state["analysis_run_id"] == first["run_id"]
    assert state["holdings"] == initial["holdings"]
    assert state["run_log"][-1]["status"] == "FAILED"
    assert any(g["field"] == "run_attempt" for g in state["data_gaps"])
    assert pipeline.validate(repo)["analytical_status"] == "FAILED"


def test_initial_invalid_input_still_has_diagnostics_workbook(repo):
    result = pipeline.run(repo, b'{"no": "schema"}')
    assert result["status"] == "FAILED"
    state = read_json(repo, pipeline.STATE)
    assert state["holdings"] == []
    assert state["summary"]["total_assets"] is None
    assert len(state["run_log"]) == 1
    pipeline.validate(repo)


def test_reconciliation_is_partial_and_missing_model_never_complete(repo, bundle):
    bundle["closing"]["cash"] = "15219"
    result = run(repo, bundle)
    assert result["status"] == "PARTIAL"
    state = read_json(repo, pipeline.STATE)
    assert state["summary"]["sizing_status"] == "BLOCKED"
    assert all(c["outcome"] == "BLOCKED_MODEL" for c in state["coverage"])


def test_replay_has_no_network_and_preserves_workspace(repo, bundle, tmp_path, monkeypatch):
    first = run(repo, bundle)
    def fail(*args, **kwargs):
        raise AssertionError("network forbidden")
    monkeypatch.setattr(socket, "create_connection", fail)
    monkeypatch.setattr(socket.socket, "connect", fail)
    original = (repo / pipeline.STATE).read_bytes()
    target = tmp_path / "replayed.xlsx"
    assert pipeline.replay(repo, first["run_id"], target)["offline"]
    validate_workbook(target, read_json(repo, pipeline.STATE))
    assert (repo / pipeline.STATE).read_bytes() == original
    assert not git(repo, "status", "--porcelain").stdout


def test_locked_workbook_rolls_back_and_recovers(repo, bundle, monkeypatch):
    first = run(repo, bundle)
    old_commit = storage.head(repo)
    old_workbook = committed(repo, pipeline.WORKBOOK)
    bundle["prices"][0]["close"] = "126"
    bundle["sources"][3]["source_id"] = "demo.prices.v2"
    for price in bundle["prices"]:
        price["source_id"] = "demo.prices.v2"
    actual = storage.atomic_write
    def locked(path, data):
        if path == repo / pipeline.WORKBOOK:
            raise PermissionError("Excel is locked")
        return actual(path, data)
    monkeypatch.setattr(storage, "atomic_write", locked)
    with pytest.raises(storage.PublicationError, match="recover"):
        run(repo, bundle)
    assert storage.head(repo) == old_commit
    assert committed(repo, pipeline.WORKBOOK) == old_workbook
    assert (repo / ".runtime/pending.json").exists()
    assert read_json(repo, pipeline.STATE)["published_run_id"] == first["run_id"]
    monkeypatch.setattr(storage, "atomic_write", actual)
    result = pipeline.recover(repo)
    assert result["status"] == "SUCCESS"
    state = read_json(repo, pipeline.STATE)
    assert Decimal(state["holdings"][0]["price"]) == 126
    assert any(l["status"] == "PUBLICATION_FAILED" for l in state["run_log"])
    assert not (repo / ".runtime/pending.json").exists()
    pipeline.validate(repo)


def test_commit_failure_recovers_and_does_not_ingest_twice(repo, bundle, monkeypatch):
    run(repo, bundle)
    actual = storage.git
    def broken(root, *args, **kwargs):
        if "commit" in args:
            raise OSError("simulated interruption before commit")
        return actual(root, *args, **kwargs)
    monkeypatch.setattr(storage, "git", broken)
    with pytest.raises(storage.PublicationError):
        run(repo, bundle)
    monkeypatch.setattr(storage, "git", actual)
    result = pipeline.recover(repo)
    state = read_json(repo, pipeline.STATE)
    assert result["status"] == "NO_MATERIAL_CHANGE"
    assert len(state["transactions"]) == 5
    assert not git(repo, "status", "--porcelain").stdout


def test_concurrent_writer_queued_and_logged(repo, bundle):
    run(repo, bundle)
    original = (repo / pipeline.WORKBOOK).read_bytes()
    with writer_lock(repo):
        with pytest.raises(storage.LockBusy):
            run(repo, bundle)
    assert (repo / pipeline.WORKBOOK).read_bytes() == original
    run(repo, bundle)
    assert any(r["status"] == "LOCK_BUSY" for r in read_json(repo, pipeline.STATE)["run_log"])


def make_note_workbook(repo, tmp_path, *, confirmed=True, field="note", value="Keep my thesis"):
    path = tmp_path / "user_notes.xlsx"
    wb = load_workbook(repo / pipeline.WORKBOOK)
    record = {"override_id": "note.demo.1", "ticker": "DEMO-BANK", "field": field, "value": value, "author": "test user", "reason": "explicit note", "effective_date": "2026-10-09", "expires_at": None, "confirmed": confirmed, "supersedes": None}
    wb["Manual_Inputs"].append([record[k] for k in NOTE_COLUMNS])
    wb.save(path)
    wb.close()
    return path


def test_explicit_notes_survive_export_and_are_idempotent(repo, bundle, tmp_path):
    run(repo, bundle)
    path = make_note_workbook(repo, tmp_path)
    pipeline.import_overrides(repo, path)
    run(repo, bundle)
    initial = read_json(repo, pipeline.STATE)
    pipeline.import_overrides(repo, path)
    state = read_json(repo, pipeline.STATE)
    assert len(state["notes"]) == 1
    assert state["notes"] == initial["notes"]
    assert state["notes"][0]["value"] == "Keep my thesis"
    assert state["notes"][0]["import_source_id"].startswith("manual.")
    assert pipeline.validate(repo)["integrity"] == "PASS"


@pytest.mark.parametrize("confirmed,field", [(False, "note"), (True, "price")])
def test_unconfirmed_and_numeric_overrides_are_logged_failures(repo, bundle, tmp_path, confirmed, field):
    run(repo, bundle)
    before = read_json(repo, pipeline.STATE)
    path = make_note_workbook(repo, tmp_path, confirmed=confirmed, field=field)
    result = pipeline.import_overrides(repo, path)
    state = read_json(repo, pipeline.STATE)
    assert result["status"] == "FAILED"
    assert state["notes"] == []
    assert state["holdings"] == before["holdings"]
    assert state["run_log"][-1]["workflow"] == "import-overrides"


def test_formula_like_note_is_literal_after_export(repo, bundle, tmp_path):
    run(repo, bundle)
    path = make_note_workbook(repo, tmp_path, value="=HYPERLINK(\"https://example.invalid\",\"note\")")
    pipeline.import_overrides(repo, path)
    wb = load_workbook(repo / pipeline.WORKBOOK)
    assert all(c.data_type != "f" for row in wb["Notes"] for c in row)
    wb.close()


def test_existing_transaction_or_source_cannot_be_rewritten(repo, bundle):
    run(repo, bundle)
    bundle["transactions"][1]["price"] = "101"
    assert run(repo, bundle)["status"] == "FAILED"
    bundle["transactions"][1]["price"] = "100"
    bundle["sources"][0]["title"] = "silently revised"
    assert run(repo, bundle)["status"] == "FAILED"


def test_manual_edits_not_overwritten(repo, bundle):
    run(repo, bundle)
    workbook = repo / pipeline.WORKBOOK
    workbook.write_bytes(b"manual edit, preserve me")
    with pytest.raises(storage.PublicationError, match="clean"):
        run(repo, bundle)
    assert workbook.read_bytes() == b"manual edit, preserve me"


def test_push_failure_preserves_local_revision(repo, bundle):
    run(repo, bundle)
    sha = storage.head(repo)
    with pytest.raises(storage.PublicationError, match="local commit retained"):
        pipeline.push(repo, "missing-remote")
    assert storage.head(repo) == sha
    assert list((repo / ".runtime/push-log").glob("*.json"))


def test_cli_reports_failed_run_exit_status(repo, tmp_path):
    path = tmp_path / "bad.json"
    path.write_text("{}")
    assert main(["--root", str(repo), "import-portfolio", "--file", str(path)]) == 1
    assert main(["--root", str(repo), "validate"]) == 0


def test_source_id_cannot_silently_change_facts(repo, bundle):
    run(repo, bundle)
    initial = read_json(repo, pipeline.STATE)
    bundle["prices"][0]["close"] = "126"
    result = run(repo, bundle)
    assert result["status"] == "FAILED"
    assert "source ID is immutable" in result["error"]
    assert read_json(repo, pipeline.STATE)["holdings"] == initial["holdings"]


def test_note_import_recovers_its_explicit_request(repo, bundle, tmp_path, monkeypatch):
    run(repo, bundle)
    path = make_note_workbook(repo, tmp_path)
    actual = storage.atomic_write
    def locked(target, data):
        if target == repo / pipeline.WORKBOOK:
            raise PermissionError("Excel locked during note import")
        return actual(target, data)
    monkeypatch.setattr(storage, "atomic_write", locked)
    with pytest.raises(storage.PublicationError):
        pipeline.import_overrides(repo, path)
    assert read_json(repo, pipeline.STATE)["notes"] == []
    monkeypatch.setattr(storage, "atomic_write", actual)
    pipeline.recover(repo)
    state = read_json(repo, pipeline.STATE)
    assert state["notes"][0]["value"] == "Keep my thesis"
    assert len(state["transactions"]) == 5


def test_post_commit_interruption_finalizes_without_double_ingestion(repo, bundle, monkeypatch):
    run(repo, bundle)
    actual = storage.atomic_write
    def interrupted(path, data):
        if path.parent == repo / ".runtime/publications":
            raise OSError("crash after Git commit")
        return actual(path, data)
    monkeypatch.setattr(storage, "atomic_write", interrupted)
    with pytest.raises(storage.PublicationError):
        run(repo, bundle)
    published = storage.head(repo)
    monkeypatch.setattr(storage, "atomic_write", actual)
    result = pipeline.recover(repo)
    assert result["status"] == "RECOVERED_COMMITTED"
    assert storage.head(repo) == published
    assert len(read_json(repo, pipeline.STATE)["transactions"]) == 5
    pipeline.validate(repo)


def test_tampered_workbook_does_not_pass_parity(repo, bundle, tmp_path):
    run(repo, bundle)
    wb = load_workbook(repo / pipeline.WORKBOOK)
    wb["Holdings"]["H2"] = 1
    target = tmp_path / "tampered.xlsx"
    wb.save(target)
    wb.close()
    with pytest.raises(ValueError, match="numeric parity"):
        validate_workbook(target, read_json(repo, pipeline.STATE))


def test_no_implicit_workbook_import(repo, bundle, tmp_path):
    run(repo, bundle)
    path = make_note_workbook(repo, tmp_path)
    wb = load_workbook(path)
    wb["Holdings"]["D2"] = 9999
    wb.save(path)
    wb.close()
    pipeline.import_overrides(repo, path)
    assert read_json(repo, pipeline.STATE)["holdings"][0]["quantity"] == "80"


def test_revision_retains_unchanged_record_lineage(repo, bundle, tmp_path):
    result = run(repo, bundle)
    before = read_json(repo, pipeline.STATE)
    pipeline.import_overrides(repo, make_note_workbook(repo, tmp_path))
    after = read_json(repo, pipeline.STATE)
    assert after["holdings"] == before["holdings"]
    assert all(h["run_id"] == result["run_id"] for h in after["holdings"])


def test_recovery_will_not_overwrite_new_manual_edits(repo, bundle, monkeypatch):
    run(repo, bundle)
    actual = storage.atomic_write
    def locked(path, data):
        if path == repo / pipeline.WORKBOOK:
            raise PermissionError("Excel locked")
        return actual(path, data)
    monkeypatch.setattr(storage, "atomic_write", locked)
    with pytest.raises(storage.PublicationError):
        run(repo, bundle)
    monkeypatch.setattr(storage, "atomic_write", actual)
    (repo / pipeline.STATE).write_bytes(b"new manual edit")
    with pytest.raises(storage.PublicationError, match="manual edit"):
        pipeline.recover(repo)
    assert (repo / pipeline.STATE).read_bytes() == b"new manual edit"
