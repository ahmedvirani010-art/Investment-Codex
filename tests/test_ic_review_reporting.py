"""Phase 3: material risk invalidation and independent Excel type/format checks."""
from copy import deepcopy
from decimal import Decimal
import json

import pytest
from openpyxl import load_workbook
from investment_workspace import pipeline
from investment_workspace.decisions import basis,calculate,BASIS_VERSION
from investment_workspace.schemas import Bundle
from investment_workspace.storage import dumps,read_json,committed
from investment_workspace.excel import validate_workbook


@pytest.fixture
def approved(repo):
    data=json.loads(pipeline.fixture_bytes(completion=True))
    initial=pipeline.run(repo,dumps(data));assert initial['status']!='FAILED',initial
    seed=read_json(repo,pipeline.STATE);p=data['integration']['proposals'][0]
    source=deepcopy(data['sources'][-1]);source['source_id']='synthetic.phase3.approval';data['sources'].append(source)
    record={'record_id':'phase3.approval','ticker':p['ticker'],'effective_date':'2026-10-09','valid_until':'2026-10-12',
      'source_ids':[source['source_id']],'proposal_id':p['record_id'],'decision':'APPROVE','decided_by':'Synthetic owner',
      'decided_at':'2026-10-09T00:00:00+05:00','reason':'Synthetic reviewed valuation and portfolio risk','reviewed_basis_hash':basis(seed,p['ticker'])}
    data['decisions']=[record]
    run=pipeline.run(repo,dumps(data));assert run['status']!='FAILED',run
    state=read_json(repo,pipeline.STATE);assert state['decisions']['rows'][0]['review_status']=='CURRENT'
    return repo,data,state,run


def test_factor_revision_within_limits_marks_approval_for_review(approved):
    repo,data,before,run=approved
    frozen=committed(repo,'runs/'+run['run_id']+'/snapshot.json')
    source=deepcopy(data['sources'][-1]);source['source_id']='synthetic.phase3.factor';data['sources'].append(source)
    factor=data['integration']['factors'][0]
    factor.update(record_id='phase3.factor',supersedes=factor['record_id'],source_ids=[source['source_id']],factors=['OIL'])
    result=pipeline.run(repo,dumps(data));assert result['status']!='FAILED',result
    after=read_json(repo,pipeline.STATE);row=after['decisions']['rows'][0]
    assert row['review_status']=='NEEDS_REVIEW' and row['basis_version']==BASIS_VERSION
    assert row['reviewed_basis_hash']==before['decisions']['rows'][0]['reviewed_basis_hash']
    assert row['basis_hash']!=before['decisions']['rows'][0]['basis_hash']
    assert after['decisions']['history']==before['decisions']['history']
    assert row['orders_created'] is False and after['transactions']==before['transactions']
    oil=next(r for r in after['integration']['risk'] if r.get('row_kind')=='FACTOR' and r['name']=='OIL')
    assert Decimal(oil['weight_total_sleeve'])>Decimal('.43') and oil['limit_exceeded'] is False
    assert committed(repo,'runs/'+run['run_id']+'/snapshot.json')==frozen
    assert pipeline.validate(repo)['workbook_parity']=='PASS'


@pytest.mark.parametrize('change',['exposure','factor_version','concentration','correlation','cash','other_holding','candidate_stop','breach'])
def test_material_review_dependencies_change_hash_and_status(approved,change):
    _,data,before,_=approved;state=deepcopy(before);ticker=data['decisions'][0]['ticker']
    if change=='exposure': state['integration']['risk'][0]['weight_total_sleeve']='.4'
    elif change=='factor_version': state['integration']['input']['factors'][0]['source_ids']=['synthetic.new-version']
    elif change=='concentration': state['summary']['hhi']='.8'
    elif change=='correlation': state['integration']['correlations'][0]['correlation']='.9'
    elif change=='cash': state['summary']['cash']='10000'
    elif change=='other_holding': next(h for h in state['holdings'] if h['ticker']!=ticker)['quantity']='999'
    elif change=='candidate_stop': state['integration']['input']['candidates'][0]['stop_price']='1'
    else: next(r for r in state['integration']['risk'] if r.get('row_kind')=='SUMMARY')['violations']=['FACTOR:OIL']
    assert basis(state,ticker)!=basis(before,ticker)
    calculate(Bundle.model_validate(data),state,before,'review.changed','hash')
    assert state['decisions']['rows'][0]['review_status']=='NEEDS_REVIEW'
    assert state['decisions']['history']==before['decisions']['history']


def test_metadata_and_list_order_do_not_invalidate_or_rewrite_decision(approved):
    repo,data,before,_=approved;state=deepcopy(before);ticker=data['decisions'][0]['ticker']
    state['notes'].append({'value':'new user note'})
    state['holdings'].reverse();state['integration']['input']['factors'].reverse()
    state['integration']['correlations'].reverse();state['integration']['risk'].reverse()
    for row in state['integration']['risk']:
        row.update(run_id='another-run',upstream_run_ids=['another-upstream'],input_hash='changed',code_hash='changed')
    assert basis(state,ticker)==basis(before,ticker)
    calculate(Bundle.model_validate(data),state,before,'unchanged','hash')
    assert state['decisions']['rows']==before['decisions']['rows']
    rerun=pipeline.run(repo,dumps(data));assert rerun['status']!='FAILED',rerun
    assert read_json(repo,pipeline.STATE)['decisions']['rows']==before['decisions']['rows']


def test_legacy_review_hash_requires_new_review(approved):
    _,data,state,_=approved;b=Bundle.model_validate(data)
    b.decisions[0].reviewed_basis_hash='0'*64
    calculate(b,state,{},'legacy','hash')
    assert state['decisions']['rows'][0]['review_status']=='NEEDS_REVIEW'
    assert state['decisions']['rows'][0]['reviewed_basis_hash']=='0'*64


MONEY=('interest_income','funding_cost','opening_equity','closing_equity','oci','dividend_capacity',
       'book_value_per_share','receivable_increase','closing_receivables','working_capital_increase','depreciation')


def test_every_engine_output_is_numeric_with_semantic_formats(approved):
    repo,_,state,_=approved;wb=load_workbook(repo/pipeline.WORKBOOK)
    ws=wb['Forward_Bridge'];headers={c.value:c.column for c in ws[1]};seen=set()
    for index,row in enumerate(state['research']['forecasts'],2):
        for field in MONEY+('sustainable_roe',):
            if row.get(field) is None: continue
            cell=ws.cell(index,headers[field]);seen.add(field)
            assert cell.data_type=='n' and isinstance(cell.value,(int,float))
            assert abs(Decimal(str(cell.value))-Decimal(row[field]))<Decimal('.00000001')
            assert cell.number_format==('0.00%' if field=='sustainable_roe' else '#,##0.00;[Red](#,##0.00)')
    assert seen==set(MONEY)|{'sustainable_roe'}
    wb.close();assert pipeline.validate(repo)['workbook_parity']=='PASS'


@pytest.mark.parametrize('corruption',['text','format'])
def test_new_workbook_rejects_type_and_format_corruption(approved,tmp_path,corruption):
    repo,_,state,_=approved;wb=load_workbook(repo/pipeline.WORKBOOK)
    ws=wb['Forward_Bridge'];cols={c.value:c.column for c in ws[1]};cell=ws.cell(2,cols['interest_income'])
    if corruption=='text': cell.value=str(cell.value)
    else: cell.number_format='General'
    target=tmp_path/'corrupted.xlsx';wb.save(target);wb.close()
    with pytest.raises(ValueError,match='parity|type/format'): validate_workbook(target,state)


def test_legacy_text_bridge_remains_readable_but_wrong_values_fail(approved,tmp_path):
    repo,_,state,_=approved;wb=load_workbook(repo/pipeline.WORKBOOK);wb.properties.version=None
    ws=wb['Forward_Bridge'];cols={c.value:c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        for field in MONEY+('sustainable_roe',):
            cell=row[cols[field]-1]
            if cell.data_type=='n': cell.value=str(cell.value);cell.number_format='General'
    target=tmp_path/'legacy.xlsx';wb.save(target)
    validate_workbook(target,state)
    ws.cell(2,cols['interest_income']).value='99999';wb.save(target);wb.close()
    with pytest.raises(ValueError,match='parity'):validate_workbook(target,state)
