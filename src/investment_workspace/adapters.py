"""Documented local adapters. No undocumented source API or implicit downloads."""

import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook

from .schemas import MarketBar


HISTORY_COLUMNS = ["session_date", "open", "high", "low", "close", "volume", "traded_value"]


def read_history_csv(path: Path):
    """Strict raw OHLCV export; provenance/calendar/actions arrive in a JSON packet."""
    with path.open(newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        if reader.fieldnames != HISTORY_COLUMNS:
            raise ValueError("CSV headers must exactly match: " + ",".join(HISTORY_COLUMNS))
        rows = []
        for row in reader:
            if None in row or any(v is None for v in row.values()):
                raise ValueError("CSV row shape differs from header")
            row["traded_value"] = row["traded_value"] or None
            rows.append(MarketBar.model_validate(row).model_dump(mode="json"))
        if not rows:
            raise ValueError("CSV contains no bars")
        return rows


def inspect_datastore(path: Path):
    """Read-only inventory; this never certifies legacy mappings or executes formulas."""
    raw = path.read_bytes()
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        for sheet in wb:
            rows = sheet.iter_rows()
            first = next(rows, [])
            headers = [str(c.value) if c.value is not None else None for c in first]
            formulas = sum(c.data_type == "f" for row in rows for c in row)
            sheets.append({"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                           "headers": headers, "formula_count": formulas})
        return {"sha256": hashlib.sha256(raw).hexdigest(), "sheets": sheets,
                "compatibility_status": "INVENTORY_ONLY_MAPPING_AND_PARITY_REQUIRED", "modified": False}
    finally:
        wb.close()
