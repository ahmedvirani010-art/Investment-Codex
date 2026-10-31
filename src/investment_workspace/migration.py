"""Explicit spreadsheet mappings, immutable originals and numerical parity."""

from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from .runtime import current, json_values
from .storage import digest, dumps


def canonical_rows(snapshot, collection, mapping=None):
    if collection=='history' and mapping and mapping.price_basis=='SPLIT_ADJUSTED':
        from .timing import adjusted_bars
        from .schemas import MarketHistory
        cutoff=date.fromisoformat(snapshot['summary']['market_cutoff'])
        return [{'ticker':history['ticker'],**json_values(bar)} for history in snapshot.get('integration',{}).get('input',{}).get('histories',[])
                if history.get('role','SECURITY')=='SECURITY' for bar in adjusted_bars(MarketHistory.model_validate(history),cutoff)]
    if collection=='history':
        return [{'ticker':history['ticker'],**bar} for history in snapshot.get('integration',{}).get('input',{}).get('histories',[])
                if history.get('role','SECURITY')=='SECURITY' for bar in history['bars']]
    module = 'integration' if collection in {'timing', 'mean_reversion'} else 'research'
    return snapshot.get(module, {}).get(collection, [])


def expected_rows(snapshot, mapping, strict=True):
    records = canonical_rows(snapshot, mapping.collection, mapping)
    if not records:
        if strict:
            raise ValueError('no canonical records for mapping: ' + mapping.collection)
        return []
    for row in records:
        if strict and any(field not in row for field in mapping.columns):
            raise ValueError('mapped canonical field missing')
    return [{field: row.get(field) for field in mapping.columns} for row in records]


def compare_workbook(raw, migration, snapshot):
    if digest(raw) != migration.workbook_sha256:
        raise ValueError('original workbook hash mismatch')
    formula = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    values = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    checks = []
    try:
        missing=set(values.sheetnames)-{m.sheet for m in migration.mappings}-set(migration.ignored_sheets)
        if missing or set(migration.ignored_sheets)-set(values.sheetnames):
            raise ValueError('unmapped legacy sheets need explicit reasons: '+', '.join(sorted(missing)))
        for mapping in migration.mappings:
            if mapping.sheet not in values:
                raise ValueError('mapped legacy sheet missing: ' + mapping.sheet)
            actual = list(values[mapping.sheet].values)
            formulas = list(formula[mapping.sheet].values)
            if len(actual) < mapping.header_row:
                raise ValueError('mapped header row missing')
            headers = list(actual[mapping.header_row-1])
            if any(headers.count(c) != 1 for c in mapping.columns.values()):
                raise ValueError('unmapped or duplicate legacy column')
            unaccounted=[h for h in headers if h is not None and h not in set(mapping.columns.values())|set(mapping.ignored_columns)]
            if unaccounted:
                raise ValueError('unmapped legacy columns need explicit reasons: '+', '.join(map(str,unaccounted)))
            if set(mapping.ignored_columns)-set(headers):
                raise ValueError('ignored legacy column not found')
            indices = {field: headers.index(col) for field, col in mapping.columns.items()}
            legacy = []
            for index, row in enumerate(actual[mapping.header_row:], mapping.header_row):
                if not any(v is not None for v in row):
                    continue
                extracted = {}
                for field, col in indices.items():
                    v = row[col] if col < len(row) else None
                    f = formulas[index][col] if col < len(formulas[index]) else None
                    if isinstance(f, str) and f.startswith('=') and v is None:
                        raise ValueError('formula has no cached result; export verified values before migration')
                    if field in {'session_date','effective_date','valid_until','period_start','period_end'}:
                        if isinstance(v,datetime): v=v.date().isoformat()
                        elif isinstance(v,date): v=v.isoformat()
                    extracted[field] = v
                legacy.append(extracted)
            expected = expected_rows(snapshot, mapping)
            def indexed(rows):
                result = {}
                for row in rows:
                    key = tuple(str(row[k]) for k in mapping.key_fields)
                    if key in result:
                        raise ValueError('duplicate mapped row key')
                    result[key] = row
                return result
            left, right = indexed(legacy), indexed(expected)
            if left.keys() != right.keys():
                raise ValueError('legacy/canonical row coverage mismatch')
            numeric_totals = {}
            for key, target in right.items():
                for field, wanted in target.items():
                    seen = left[key][field]
                    if field in mapping.key_fields or isinstance(wanted, (bool, list, dict)):
                        if seen != wanted:
                            raise ValueError('legacy value mismatch: ' + field)
                        continue
                    try:
                        n = Decimal(str(wanted))
                        numeric = n.is_finite()
                    except InvalidOperation:
                        numeric = False
                    if wanted is None:
                        if seen not in (None, 'UNKNOWN'):
                            raise ValueError('unknown legacy value became a number')
                    elif numeric:
                        try:
                            number = Decimal(str(seen))
                        except InvalidOperation as exc:
                            raise ValueError('legacy numeric value missing: ' + field) from exc
                        if not number.is_finite() or abs(number-n) > max(migration.absolute_tolerance, abs(n)*migration.relative_tolerance):
                            raise ValueError('legacy numeric parity mismatch: ' + field)
                        numeric_totals[field] = numeric_totals.get(field, Decimal(0)) + n
                    elif seen != wanted:
                        raise ValueError('legacy text mismatch: ' + field)
            checks.append({'sheet': mapping.sheet, 'collection': mapping.collection, 'row_count': len(expected),
                           'totals': {k: str(v) for k, v in numeric_totals.items()}, 'parity': 'PASS',
                           'ignored_columns':mapping.ignored_columns,
                           'formula_count':sum(isinstance(v,str) and v.startswith('=') for row in formulas[mapping.header_row:] for v in row)})
    finally:
        formula.close()
        values.close()
    history_complete=any(m.collection=='history' and {'ticker','session_date','close'}<=m.columns.keys() for m in migration.mappings)
    signals_complete=any(m.collection in {'timing','mean_reversion'} and 'ticker' in m.columns for m in migration.mappings)
    certified=migration.kind!='LEGACY_DATASTORE' or history_complete and signals_complete and migration.methodology_reviewed
    return {'status': 'VERIFIED' if certified else 'PARTIAL_MAPPING', 'workbook_sha256': migration.workbook_sha256, 'checks': checks,
            'baseline_hash': digest(dumps(checks)), 'record_id': migration.record_id}


def verify_migrations(bundle, snapshot, previous):
    ctx = current()
    old = {r['payload']['record_id']: r for r in previous.get('migration_records', [])}
    records = []
    for migration in bundle.migrations:
        if migration.effective_date>bundle.market_cutoff:
            raise ValueError('post-cutoff migration record')
        payload = migration.model_dump(mode='json')
        if not all(s.evidence_label == 'CONFIRMED' for s in bundle.sources if s.source_id in migration.source_ids) or not set(migration.source_ids) <= {s.source_id for s in bundle.sources}:
            raise ValueError('migration sources must be registered and confirmed')
        hashes={s.document_sha256 for s in bundle.sources if s.source_id in migration.source_ids and s.document_sha256}
        if hashes and migration.workbook_sha256 not in hashes:
            raise ValueError('migration source document hash mismatch')
        raw = ctx.blob(migration.workbook_sha256, '.xlsx') if ctx else None
        if raw is None:
            raise ValueError('migration original not retained as immutable evidence')
        if digest(raw) != migration.workbook_sha256:
            raise ValueError('migration evidence corrupt')
        if migration.supersedes:
            target=next((m for m in bundle.migrations if m.record_id==migration.supersedes),None)
            if not target or target.kind!=migration.kind or target.effective_date>migration.effective_date:
                raise ValueError('invalid migration supersession')
        if migration.record_id in old:
            if old[migration.record_id]['payload'] != payload:
                raise ValueError('migration record immutable')
            records.append(old[migration.record_id])
        else:
            if migration.kind == 'LEGACY_DATASTORE' and not any(m.collection in {'timing', 'mean_reversion'} for m in migration.mappings):
                raise ValueError('datastore migration must reconcile timing/reversion records')
            records.append({'payload': payload, 'validation': compare_workbook(raw, migration, snapshot)})
    if set(old) - {m.record_id for m in bundle.migrations}:
        raise ValueError('accepted migration evidence cannot be omitted')
    snapshot['migration_records'] = records
    return bool(active_legacy(records))


def export_compatibility(path, snapshot, blocked=False):
    wb = Workbook()
    wb.remove(wb.active)
    seen = set()
    from .schemas import SheetMapping
    for accepted in active_legacy(snapshot.get('migration_records', [])):
        if accepted['payload']['kind'] != 'LEGACY_DATASTORE':
            continue
        for raw_mapping in accepted['payload']['mappings']:
            mapping = SheetMapping.model_validate(raw_mapping)
            if mapping.sheet in seen:
                raise ValueError('conflicting compatibility sheet mappings')
            seen.add(mapping.sheet)
            ws = wb.create_sheet(mapping.sheet)
            for _ in range(mapping.header_row-1):
                ws.append([None])
            ws.append(list(mapping.columns.values()))
            for row in expected_rows(snapshot, mapping, strict=False):
                values = []
                for field in mapping.columns:
                    v = row[field]
                    try:
                        n = Decimal(str(v))
                        if field not in mapping.key_fields and n.is_finite():
                            v = float(n)
                    except (InvalidOperation, ValueError):
                        pass
                    values.append('UNKNOWN' if v is None else v)
                ws.append(values)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            ws.freeze_panes = 'A' + str(mapping.header_row+1)
            ws.auto_filter.ref = f'A{mapping.header_row}:{ws.cell(ws.max_row,ws.max_column).coordinate}'
    if not seen and not blocked:
        raise ValueError('no verified compatibility mapping')
    provenance=wb.create_sheet('_Workspace_Provenance')
    provenance.append(['field','value'])
    for key,value in [('status','NOT_CURRENT_COMPATIBILITY' if blocked else 'DERIVED_FROM_CANONICAL'),
                      ('run_id',snapshot['published_run_id']),('market_cutoff',snapshot['summary']['market_cutoff']),
                      ('evaluation_as_of',snapshot.get('evaluation_as_of',snapshot['summary']['market_cutoff'])),
                      ('original_evidence','Immutable originals and prior generated versions remain in Git history')]:
        provenance.append([key,value])
    provenance.freeze_panes='A2'
    provenance.column_dimensions['A'].width=25
    provenance.column_dimensions['B'].width=65
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def active_legacy(records):
    superseded={r['payload'].get('supersedes') for r in records}
    return [r for r in records if r['payload']['kind']=='LEGACY_DATASTORE' and r['payload']['record_id'] not in superseded
            and r['validation']['status']=='VERIFIED']
