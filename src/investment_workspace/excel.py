"""Computed workbook cells; no Excel recalculation or implicit numerical overrides."""

import json
import math
import hashlib
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .schemas import NoteOverride
from .research import SHEETS
from .workflows import SHEETS as INTEGRATION_SHEETS
from .wealth import SHEETS as WEALTH_SHEETS

NOTE_COLUMNS = ["override_id", "ticker", "field", "value", "author", "reason", "effective_date", "expires_at", "confirmed", "supersedes"]
HOLDING_COLUMNS = ["record_id", "ticker", "sector", "quantity", "cost", "price", "price_date", "market_value", "unrealized_gain", "unrealized_return_on_cost", "equity_weight", "total_sleeve_weight", "freshness", "reconciliation_state", "run_id", "effective_date", "source_ids"]
FUTURE_SHEETS = []
NUMERIC_COLUMNS = {"quantity", "cost", "price", "market_value", "unrealized_gain", "unrealized_return_on_cost", "equity_weight", "total_sleeve_weight", "cash_after", "realized_gain", "amount", "fees", "taxes", "ratio", "present_fair_value", "horizon_exit_price", "horizon_dividends", "forward_eps", "normalized_eps", "reported_eps", "normalized_pat", "sustainable_dps", "entry_ceiling", "mos_entry_ceiling", "hurdle_entry_ceiling", "hurdle", "gross_return", "net_return", "annualized_return", "margin_of_safety", "dps", "eligible_quantity", "gross_entitlement", "net_entitlement", "gross_received", "net_received", "gross_pending", "payout_crosscheck", "horizon_years", "price_return", "benchmark_return", "relative_performance", "rs_momentum", "short_sma", "long_sma", "sma_slope", "extension", "z_score", "average_traded_value", "rs_percentile", "price_breadth", "rs_breadth", "average_relative_performance", "weight_total_sleeve", "max_weight", "scenario_assets", "income_return", "total_return", "correlation", "desired_entry", "stop_price", "approach_fraction", "balance_pkr", "spendable_now", "nominal_balance", "real_balance", "accessible_balance", "locked_balance", "growth", "net_inflows", "net_spending_required", "net_withdrawn", "funding_gap", "cumulative_funding_gap", "inflation_index", "ending_nominal_balance", "ending_real_balance", "holding_period_return", "revenue", "ebitda", "pretax", "tax_charge", "pat", "eps", "fcfe", "shares", "ledger_quantity", "ledger_cost", "quantity_difference", "cost_difference", "cash_difference"}

NUMERIC_COLUMNS.update({"starting_wealth_pkr","monthly_spending_real","annual_inflation",
                        "ending_accessible_balance","ending_locked_balance","left_common_return","right_common_return"})


def cell_value(value, column=""):
    if value is None:
        return "UNKNOWN"
    if isinstance(value,(list,dict)):
        raw=json.dumps(value,sort_keys=True,ensure_ascii=False)
        if len(raw)>32000:
            return json.dumps({'status':'FULL_VALUE_IN_CANONICAL_SNAPSHOT','items':len(value),
                               'sha256':hashlib.sha256(raw.encode()).hexdigest()},sort_keys=True)
        return raw
    if column in NUMERIC_COLUMNS:
        exact=Decimal(str(value))
        result=float(exact)
        return result if math.isfinite(result) and (result!=0 or exact==0) else str(value)
    if isinstance(value,str) and len(value)>32000:
        return 'FULL_TEXT_IN_CANONICAL_SNAPSHOT sha256='+hashlib.sha256(value.encode()).hexdigest()
    return value


def add_sheet(wb, name, rows, columns=None):
    ws = wb.create_sheet(name)
    columns = columns or list(dict.fromkeys(key for row in rows for key in row)) or ["status", "reason"]
    ws.append(columns)
    for row in rows:
        ws.append([cell_value(row.get(col), col) for col in columns])
        # Untrusted notes/titles must stay literal; prevent spreadsheet formula injection.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for index, col in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = min(48, max(16, len(col) + 3))
        for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            cell = row[0]
            if col in {"unrealized_return_on_cost", "equity_weight", "total_sleeve_weight", "gross_return", "net_return", "annualized_return", "margin_of_safety", "hurdle", "price_return", "benchmark_return", "relative_performance", "rs_momentum", "sma_slope", "extension", "weight_total_sleeve", "max_weight", "price_breadth", "rs_breadth", "income_return", "total_return"} and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"
            elif col in NUMERIC_COLUMNS and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;[Red](#,##0.00)'
            if cell.value in ("UNKNOWN", "STALE", "EXCEPTION", "FAILED", "BLOCKED", "PUBLICATION_FAILED"):
                cell.fill = PatternFill("solid", fgColor="FFE699")
    if rows:
        table = Table(displayName=f"T_{name}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = ws.dimensions
    return ws


def export_workbook(path: Path, snapshot: dict):
    wb = Workbook()
    wb.remove(wb.active)
    dashboard = [{"metric": "workspace", "value": "SYNTHETIC DEMO — NOT AHMED'S HOLDINGS" if snapshot["summary"].get("mode") == "DEMO" else "USER IMPORT — CHECK GATES"}]
    dashboard += [{"metric": "attempt_status", "value": snapshot["attempt_status"]}, {"metric": "publication_status", "value": "LOCAL_ONLY"}, {"metric": "published_run_id", "value": snapshot["published_run_id"]}, {"metric": "analysis_run_id", "value": snapshot.get("analysis_run_id")}, {"metric": "valuation_status", "value": snapshot.get("research", {}).get("summary", {}).get("valuation_status", "BLOCKED_MODEL — no research imported")}]
    dashboard += [{"metric": key, "value": value} for key, value in snapshot["summary"].items()]
    dashboard += [{"metric": key if key.startswith("integration_") else "integration_" + key, "value": value} for key, value in snapshot.get("integration", {}).get("summary", {}).items()]
    if not snapshot.get("integration"):
        dashboard.append({"metric": "integration_status", "value": "BLOCKED_INPUT — no integration imported"})
    dashboard += [{"metric":"wealth_"+key,"value":value} for key,value in snapshot.get("wealth",{}).get("summary",{}).items()]
    add_sheet(wb, "Dashboard", dashboard, ["metric", "value"])
    add_sheet(wb, "Holdings", snapshot["holdings"], HOLDING_COLUMNS)
    add_sheet(wb, "Transactions", snapshot["transactions"])
    add_sheet(wb, "Cash", [{k: snapshot["summary"].get(k) for k in ("cash", "ledger_cash", "cash_difference", "income_received", "realized_gain")}])
    add_sheet(wb, "Sources", snapshot["sources"])
    add_sheet(wb, "Coverage", snapshot["coverage"])
    add_sheet(wb, "Data_Gaps", snapshot["data_gaps"])
    add_sheet(wb, "Reconciliation", snapshot["reconciliation"])
    add_sheet(wb, "Run_Log", snapshot["run_log"])
    add_sheet(wb, "Change_Log", snapshot["changes"])
    add_sheet(wb, "Usage", snapshot["usage"])
    add_sheet(wb, "Notes", snapshot["notes"])
    add_sheet(wb, "Assumptions", snapshot.get("assumptions", []))
    add_sheet(wb, "Manual_Inputs", [], NOTE_COLUMNS)
    for name, key in SHEETS.items():
        add_sheet(wb, name, snapshot.get("research", {}).get(key, []))
    for name, key in INTEGRATION_SHEETS.items():
        add_sheet(wb, name, snapshot.get("integration", {}).get(key, []))
    for name,key in WEALTH_SHEETS.items():
        add_sheet(wb,name,snapshot.get('wealth',{}).get(key,[]))
    add_sheet(wb,'Financial_Facts',snapshot.get('financials',{}).get('facts',[]))
    add_sheet(wb,'Evidence_Library',snapshot.get('documents',[]))
    add_sheet(wb,'Readiness',[{'dimension':key,'value':value} for key,value in sorted(snapshot.get('readiness',{}).items())])
    for name in FUTURE_SHEETS:
        add_sheet(wb, name, [{"status": "NOT_IMPLEMENTED", "reason": "No accepted model or live data migrated; see docs/migration.md"}])
    if 'url' in [c.value for c in wb['Sources'][1]]:
        index=[c.value for c in wb['Sources'][1]].index('url')+1
        for row in wb['Sources'].iter_rows(min_row=2):
            cell=row[index-1]
            if isinstance(cell.value,str) and cell.value.startswith(('https://','http://')):
                cell.hyperlink=cell.value
                cell.style='Hyperlink'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    validate_workbook(path, snapshot)


def validate_workbook(path: Path, snapshot: dict):
    wb = load_workbook(path, data_only=False)
    try:
        expected_rows = [[cell_value(row.get(col), col) for col in HOLDING_COLUMNS] for row in snapshot["holdings"]]
        actual_rows = list(wb["Holdings"].values)
        if actual_rows[0] != tuple(HOLDING_COLUMNS) or len(actual_rows) - 1 != len(expected_rows):
            raise ValueError("workbook holdings shape differs from canonical records")
        for actual, expected in zip(actual_rows[1:], expected_rows):
            for left, right in zip(actual, expected):
                if isinstance(right, float):
                    if not isinstance(left, (float, int)) or abs(left - right) > max(1e-8, abs(right) * 1e-12):
                        raise ValueError("workbook numeric parity failed")
                elif left != right:
                    raise ValueError("workbook record parity failed")
        logs = list(wb["Run_Log"].values)
        if len(logs) - 1 != len(snapshot["run_log"]):
            raise ValueError("workbook run log parity failed")
        dashboard = dict(list(wb["Dashboard"].values)[1:])
        for key in ("cash", "total_assets", "market_cutoff", "sizing_status"):
            if dashboard.get(key) != cell_value(snapshot["summary"].get(key)):
                raise ValueError(f"workbook dashboard parity failed: {key}")
        if tuple(int(v) for v in snapshot["summary"].get("code_version","0").split(".")) >= (0,5):
            for key,value in snapshot.get("wealth",{}).get("summary",{}).items():
                if dashboard.get("wealth_"+key)!=cell_value(value):
                    raise ValueError("workbook wealth dashboard parity failed: "+key)
        if dashboard["published_run_id"] != snapshot["published_run_id"]:
            raise ValueError("workbook revision differs from dataset")
        tables = {"Transactions": snapshot["transactions"], "Cash": [{k: snapshot["summary"].get(k) for k in ("cash", "ledger_cash", "cash_difference", "income_received", "realized_gain")}], "Sources": snapshot["sources"], "Coverage": snapshot["coverage"], "Data_Gaps": snapshot["data_gaps"], "Reconciliation": snapshot["reconciliation"], "Run_Log": snapshot["run_log"], "Change_Log": snapshot["changes"], "Usage": snapshot["usage"], "Notes": snapshot["notes"], "Assumptions": snapshot.get("assumptions", [])}
        for module, sheets in (("research", SHEETS), ("integration", INTEGRATION_SHEETS), ("wealth", WEALTH_SHEETS)):
            for name, key in sheets.items():
                if name not in wb:
                    if key in snapshot.get(module,{}):
                        raise ValueError("workbook missing analytical sheet: " + name)
                    continue
                legacy = list(wb[name].values)
                placeholder = (not snapshot.get(module) and len(legacy) == 2 and legacy[0] == ("status", "reason")
                               and legacy[1][0] == "NOT_IMPLEMENTED")
                if not placeholder:
                    tables[name] = snapshot.get(module, {}).get(key, [])
        if snapshot.get('financials') and 'Financial_Facts' not in wb or snapshot.get('readiness') and 'Readiness' not in wb or snapshot.get('documents') and 'Evidence_Library' not in wb:
            raise ValueError('workbook missing evidence/readiness table')
        if 'Evidence_Library' in wb:
            tables['Evidence_Library']=snapshot.get('documents',[])
        if 'Financial_Facts' in wb:
            tables['Financial_Facts']=snapshot.get('financials',{}).get('facts',[])
        if 'Readiness' in wb:
            tables['Readiness']=[{'dimension':key,'value':value} for key,value in sorted(snapshot.get('readiness',{}).items())]
        for name, records in tables.items():
            values = list(wb[name].values)
            headers = values[0]
            expected_keys = set(key for record in records for key in record)
            if records and set(headers) != expected_keys:
                raise ValueError(f"workbook columns differ from canonical records: {name}")
            if len(values) - 1 != len(records):
                raise ValueError(f"workbook row count differs from canonical records: {name}")
            for actual, record in zip(values[1:], records):
                expected = [cell_value(record.get(key), key) for key in headers]
                for left, right in zip(actual, expected):
                    if isinstance(right, float):
                        if not isinstance(left, (float, int)) or abs(left - right) > max(1e-8, abs(right) * 1e-12):
                            raise ValueError(f"workbook numeric parity failed: {name}")
                    elif left != right:
                        raise ValueError(f"workbook record parity failed: {name}")
        for key, value in snapshot["summary"].items():
            if dashboard.get(key) != cell_value(value):
                raise ValueError(f"workbook summary parity failed: {key}")
        for sheet in wb:
            if any(c.data_type == "f" for row in sheet for c in row):
                raise ValueError("unexpected executable workbook formula")
    finally:
        wb.close()


def read_notes(path: Path) -> list[NoteOverride]:
    wb = load_workbook(path, data_only=False)
    try:
        if "Manual_Inputs" not in wb:
            raise ValueError("workbook has no Manual_Inputs sheet")
        rows = list(wb["Manual_Inputs"].values)
        if not rows or list(rows[0]) != NOTE_COLUMNS:
            raise ValueError("Manual_Inputs headers differ from the supported note contract")
        notes = []
        for values in rows[1:]:
            if not any(v is not None for v in values):
                continue
            record = dict(zip(NOTE_COLUMNS, values))
            for key in ("effective_date", "expires_at"):
                if hasattr(record[key], "date"):
                    record[key] = record[key].date()
            for key in ("expires_at", "supersedes"):
                if record[key] == "":
                    record[key] = None
            if record["confirmed"] in ("YES", "TRUE", "true"):
                record["confirmed"] = True
            notes.append(NoteOverride.model_validate(record))
        return notes
    finally:
        wb.close()
