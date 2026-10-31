import json,tempfile,io
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from investment_workspace import pipeline
from investment_workspace.storage import dumps,read_json,committed
from investment_workspace.decisions import basis
x=json.loads(pipeline.fixture_bytes(completion=True));out={}
with tempfile.TemporaryDirectory(prefix='ic-audit-') as tmp:
 root=Path(tmp)/'workspace';pipeline.run(root,dumps(x));seed=read_json(root,pipeline.STATE)
 wb=load_workbook(io.BytesIO(committed(root,pipeline.WORKBOOK)))
 sheet=wb['Forward_Bridge'];cols={c.value:c.column for c in sheet[1]};out['excel_bank_bridge']={}
 for field in ('pat','interest_income','funding_cost','opening_equity','closing_equity','sustainable_roe','dividend_capacity'):
  cell=sheet.cell(2,cols[field]);out['excel_bank_bridge'][field]={'value':cell.value,'data_type':cell.data_type,'format':cell.number_format}
 wb.close()
 def source(tag):
  s=deepcopy(x['sources'][-1]);s['source_id']='synthetic.audit.'+tag;x['sources'].append(s);return s['source_id']
 p=x['integration']['proposals'][0]
 x['decisions']=[{'record_id':'decision.audit','ticker':p['ticker'],'effective_date':'2026-10-09','valid_until':'2026-10-12','source_ids':[source('decision')],
   'proposal_id':p['record_id'],'decision':'APPROVE','decided_by':'Synthetic audit','decided_at':'2026-10-09T00:00:00+05:00','reason':'Probe risk invalidation',
   'reviewed_basis_hash':basis(seed,p['ticker'])}]
 r=pipeline.run(root,dumps(x));assert r['status']!='FAILED',r
 before=read_json(root,pipeline.STATE);f=x['integration']['factors'][0]
 f.update(record_id='audit.corrected-factor',supersedes=f['record_id'],source_ids=[source('factor')],factors=['OIL'])
 r=pipeline.run(root,dumps(x));assert r['status']!='FAILED',r
 after=read_json(root,pipeline.STATE)
 def factor_rows(state):return [{k:r.get(k) for k in ('name','weight_total_sleeve','limit_exceeded')} for r in state['integration']['risk'] if r.get('row_kind')=='FACTOR']
 out['ic_factor_change']={'before_review':before['decisions']['rows'][0]['review_status'],'after_review':after['decisions']['rows'][0]['review_status'],
 'basis_unchanged':basis(before,p['ticker'])==basis(after,p['ticker']),'before_factors':factor_rows(before),'after_factors':factor_rows(after)}
Path('/tmp/prd-ic-excel-probe.json').write_bytes(dumps(out));print(json.dumps(out,indent=2))
